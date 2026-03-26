import os
import pandas as pd
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, send_file, make_response
from sqlalchemy import create_engine, text, inspect
import sys
from io import BytesIO
from dotenv import load_dotenv
from VERSION import get_version
from weasyprint import HTML
import base64
import json
from datetime import datetime
from html import escape
import requests

# Carrega .env se existir (apenas local)
load_dotenv()

app = Flask(__name__)
# Chave secreta definida pelo usuário
app.secret_key = os.getenv('SECRET_KEY', '114211Jo@')

print("\n" + "="*40, flush=True)
print("--- DIAGNÓSTICO DE INICIALIZAÇÃO (V7 - FIX DATA BRASIL) ---", flush=True)

# Lógica de Prioridade: DATABASE_URL > Variáveis Separadas
db_url = os.getenv('DATABASE_URL')

if db_url:
    print("AVISO: Encontrada variável DATABASE_URL. Ela terá prioridade.", flush=True)
    if db_url.startswith("mysql://"):
        db_url = db_url.replace("mysql://", "mysql+pymysql://")
    print(f"      URL em uso: {db_url.split('@')[-1]}", flush=True)
else:
    print("INFO: Usando variáveis separadas (DB_HOST, DB_NAME...)", flush=True)
    DB_HOST = os.getenv('DB_HOST', 'site_nir')
    DB_USER = os.getenv('DB_USER', 'joabeoliveira')
    DB_PASS = os.getenv('DB_PASSWORD', '114211Jo')
    DB_NAME = os.getenv('DB_NAME', 'nir')
    DB_PORT = os.getenv('DB_PORT', '3306')
    
    print(f"      Host: {DB_HOST}", flush=True)
    print(f"      Banco Alvo: {DB_NAME}", flush=True)
    
    db_url = f"mysql+pymysql://{DB_USER}:{DB_PASS}@{DB_HOST}:{DB_PORT}/{DB_NAME}"

print("="*40 + "\n", flush=True)

# Variáveis globais
engine = None
db_status = False
MAX_RELATORIO_BLOCKS = int(os.getenv('MAX_RELATORIO_BLOCKS', '12'))
MAX_TABELA_ROWS = int(os.getenv('MAX_TABELA_ROWS', '100'))
WEBHOOK_TIMEOUT_SECONDS = int(os.getenv('WEBHOOK_TIMEOUT_SECONDS', '20'))
WEBHOOK_CONFIG = {
    "url": os.getenv('N8N_WEBHOOK_URL', '').strip(),
    "enabled": bool(os.getenv('N8N_WEBHOOK_URL', '').strip())
}

try:
    engine = create_engine(db_url)
    with engine.connect() as conn:
        db_atual = conn.execute(text("SELECT DATABASE()")).scalar()
        print(f">>> CONECTADO COM SUCESSO! Banco Atual: '{db_atual}' <<<", flush=True)
        db_status = True
except Exception as e:
    print(f">>> FALHA DE CONEXÃO: {e}", flush=True)

# --- MAPEAMENTO DE COLUNAS ---
DE_PARA = {
    'NUM ENF': 'num_enf', 'LEITO': 'leito', 'NOME ENFERMARIA': 'nome_enfermaria', 
    'STATUS': 'status_leito', 'AIH GERADA PARA O PACIENTE': 'aih_paciente', 
    'NÚMERO DO CNS DO PACIENTE': 'cns_paciente', 'NOME PACIENTE': 'nome_paciente', 
    'SEXO': 'sexo', 'DATA NASC': 'data_nasc', 'IDADE': 'idade', 
    'DATA INTERN': 'data_internacao', 'DATA INTERN LEITO': 'data_internacao_leito', 
    'SUSPEITA COVID': 'suspeita_covid', 'PÓS COVID': 'pos_covid', 
    'MOTIVO IMPEDIMENTO': 'motivo_impedimento', 'DATA IMPEDIMENTO': 'data_impedimento', 
    'DATA SOL. RESERVA': 'data_sol_reserva', 'PREVISÃO INTERN. RESERVA': 'previsao_intern_reserva', 
    'ACOMPANHAMENTO DATA / HORA': 'acompanhamento_data_hora', 'PRONTUÁRIO': 'prontuario', 
    'CID 10': 'cid_10', 'CÓDIGO SER': 'codigo_ser', 'PERFIL': 'perfil', 
    'BOMBA INFUSORA': 'bomba_infusora', 'SUPORTE CIRÚRGICO': 'suporte_cirurgico', 
    'SUPORTE ALIMENTAR': 'suporte_alimentar', 'MODO VENTILATÓRIO': 'modo_ventilatorio', 
    'OBSERVAÇÃO': 'observacao', 'CRÔNICO': 'cronico', 'LONGA PERMANÊNCIA': 'longa_permanencia', 
    'SITUAÇÃO / MOTIVO PERMANÊNCIA': 'situacao_motivo_permanencia', 'GESTANTE': 'gestante', 
    'INDUÇÃO AO PARTO': 'inducao_parto', 'ARBOVIROSE': 'arbovirose', 
    'VIABILIDADE PARA DIÁLISE PERITONEAL': 'viabilidade_dialise_peritoneal', 
    'INFECÇÃO ATIVA E/OU USO DE ANTIBIOTICO': 'infeccao_ativa_antibiotico', 
    'HISTÓRICO DE CIRURGIAS INTRA-ABDOMINAIS': 'historico_cirurgias_abdominais', 
    'DOENÇA NEOPLÁSICA AVANÇADA': 'doenca_neoplasica_avancada', 
    'HÉRNIA INGUINAL A REPARAR': 'hernia_inguinal_reparar', 
    'CONDIÇÕES DE ALTA COM DIÁLISE': 'condicoes_alta_dialise', 
    'ESTÁ INSERIDO NO TRS': 'inserido_no_trs'
}

# --- ENFERMARIAS DE EMERGÊNCIA ---
EMERGENCY_WARDS = [
    'CLINICA REFERENCIADA',
    'CIRURGICA REFERENCIADA',
    'CIRURGICA REFERENCIADA - FEMININA',
    'CIRURGICA REFERENCIADA - MASCULINA',
    'CLINICA REFERENCIADA - PED'
]

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    file = request.files.get('file')
    data_ref_input = request.form.get('data_referencia') 

    if not file or not data_ref_input:
        flash('Arquivo e Data são obrigatórios.', 'error')
        return redirect(url_for('index'))

    try:
        # --- CORREÇÃO DE DATA (V2) ---
        # Força dayfirst=True para evitar que 05/01 (Jan) vire 01/05 (Maio)
        try:
            dt_obj = pd.to_datetime(data_ref_input, dayfirst=True)
            data_banco = dt_obj.strftime('%Y-%m-%d')  # Formato MySQL (Ano-Mês-Dia)
            data_visual = dt_obj.strftime('%d/%m/%Y') # Formato Visual (Dia/Mês/Ano)
            print(f"Data recebida: {data_ref_input} -> Interpretada como: {data_visual}", flush=True)
        except:
            flash(f'Formato de data inválido: {data_ref_input}', 'error')
            return redirect(url_for('index'))

        # Leitura do arquivo: aceita CSV e XLS/XLSX e ignora as 2 primeiras linhas (header na 3ª linha)
        filename = (file.filename or '').lower()
        df = None
        try:
            file.seek(0)
            if filename.endswith('.xlsx'):
                # explicit engine for .xlsx — read into BytesIO to avoid ambiguity
                try:
                    file.seek(0)
                    content = file.read()
                    bio = BytesIO(content)
                    # Detect HTML-like uploads (sometimes Excel saved as webpage)
                    head = content.lstrip()[:64].lower()
                    if head.startswith(b'<') or b'<html' in head or b'<!doctype html' in head:
                        # Try to parse HTML tables as fallback
                        try:
                            dfs = pd.read_html(BytesIO(content), header=2)
                            df = dfs[0]
                            print(f"Fallback: parsed HTML-like .xlsx upload via read_html ({len(content)} bytes)", flush=True)
                        except Exception as e:
                            print(f"ERRO LEITURA XLSX (html fallback): {e}", flush=True)
                            flash('Arquivo parece ser HTML/corrompido. Salve o arquivo como .xlsx válido e tente novamente.', 'error')
                            return redirect(url_for('index'))
                    else:
                        df = pd.read_excel(bio, header=2, engine='openpyxl')
                        print(f"Leitura .xlsx via openpyxl ({len(content)} bytes)", flush=True)
                except Exception as e:
                    print(f"ERRO LEITURA XLSX: {e}", flush=True)
                    flash('Erro ao ler .xlsx: verifique se o arquivo é válido e se openpyxl está instalado.', 'error')
                    return redirect(url_for('index'))
            elif filename.endswith('.xls'):
                # explicit engine for old .xls files — use BytesIO
                try:
                    file.seek(0)
                    content = file.read()
                    bio = BytesIO(content)
                    head = content.lstrip()[:64].lower()
                    if head.startswith(b'<') or b'<html' in head or b'<!doctype html' in head:
                        try:
                            dfs = pd.read_html(BytesIO(content), header=2)
                            df = dfs[0]
                            print(f"Fallback: parsed HTML-like .xls upload via read_html ({len(content)} bytes)", flush=True)
                        except Exception as e:
                            print(f"ERRO LEITURA XLS (html fallback): {e}", flush=True)
                            flash('Arquivo parece ser HTML/corrompido. Salve o arquivo como .xls válido e tente novamente.', 'error')
                            return redirect(url_for('index'))
                    else:
                        df = pd.read_excel(bio, header=2, engine='xlrd')
                        print(f"Leitura .xls via xlrd ({len(content)} bytes)", flush=True)
                except Exception as e:
                    print(f"ERRO LEITURA XLS: {e}", flush=True)
                    flash('Erro ao ler .xls: verifique se o arquivo é válido e se xlrd (2.0.1) está instalado.', 'error')
                    return redirect(url_for('index'))
            else:
                # CSV: tenta utf-8 com vírgula, se falhar tenta latin1 com ponto-e-vírgula
                try:
                    df = pd.read_csv(file, header=2, encoding='utf-8', sep=',')
                except Exception:
                    file.seek(0)
                    df = pd.read_csv(file, header=2, encoding='latin1', sep=';')
        except Exception as e:
            print(f"ERRO LEITURA ARQUIVO (geral): {e}", flush=True)
            flash(f'Erro ao ler o arquivo: {str(e)}', 'error')
            return redirect(url_for('index'))

        # ETL
        df_banco = df.rename(columns=DE_PARA)
        cols_uteis = [c for c in df_banco.columns if c in DE_PARA.values()]
        df_banco = df_banco[cols_uteis]
        
        # Usa a data formatada e segura
        df_banco['data_referencia'] = data_banco
        
        # --- LIMPEZA BLINDADA ---
        if 'num_enf' in df_banco.columns:
            df_banco['num_enf'] = pd.to_numeric(df_banco['num_enf'], errors='coerce')
            df_banco = df_banco.dropna(subset=['num_enf'])

        for col in ['cns_paciente', 'prontuario', 'aih_paciente']:
            if col in df_banco.columns:
                df_banco[col] = df_banco[col].astype(str).str.replace(r'\.0$', '', regex=True).replace('nan', None)
        
        for col in df_banco.columns:
            if 'data' in col or 'previsao' in col:
                df_banco[col] = pd.to_datetime(df_banco[col], dayfirst=True, errors='coerce')

        # Gravação
        insp = inspect(engine)
        tabela_existe = insp.has_table('historico_ocupacao_completo')
        
        with engine.begin() as conn:
            if tabela_existe:
                # Remove dados antigos dessa data específica
                conn.execute(text(f"DELETE FROM historico_ocupacao_completo WHERE data_referencia = '{data_banco}'"))
            
            df_banco.to_sql('historico_ocupacao_completo', con=conn, if_exists='append', index=False)

        # MENSAGEM MELHORADA: Mostra a data no formato brasileiro para confirmação
        flash(f'Sucesso! {len(df_banco)} registros importados para o dia {data_visual}.', 'success')
        return redirect(url_for('index', data=data_banco))
        
    except Exception as e:
        print(f"ERRO DE UPLOAD: {e}", flush=True)
        flash(f'Erro técnico: {str(e)}', 'error')
        return redirect(url_for('index'))


@app.route('/fix_date', methods=['POST'])
def fix_date():
    """Move records from one data_referencia to another (correction tool).

    Expects form fields: from_date (errada) and to_date (correta).
    """
    from_date = request.form.get('from_date')
    to_date = request.form.get('to_date')

    if not from_date or not to_date:
        flash('Ambas as datas (errada e correta) são obrigatórias.', 'error')
        return redirect(url_for('index'))

    # Parse dates safely
    try:
        dt_from = pd.to_datetime(from_date, dayfirst=True)
        dt_to = pd.to_datetime(to_date, dayfirst=True)
        db_from = dt_from.strftime('%Y-%m-%d')
        db_to = dt_to.strftime('%Y-%m-%d')
        visual_from = dt_from.strftime('%d/%m/%Y')
        visual_to = dt_to.strftime('%d/%m/%Y')
    except Exception as e:
        flash(f'Formato de data inválido: {str(e)}', 'error')
        return redirect(url_for('index'))

    try:
        with engine.begin() as conn:
            # Count how many rows exist for from_date
            cnt = conn.execute(text("SELECT COUNT(*) FROM historico_ocupacao_completo WHERE data_referencia = :d"), {"d": db_from}).scalar()
            if cnt == 0:
                flash(f'Nenhum registro encontrado para {visual_from}. Nada foi alterado.', 'error')
                return redirect(url_for('index'))

            # Perform update: move rows to new date
            result = conn.execute(text("UPDATE historico_ocupacao_completo SET data_referencia = :to WHERE data_referencia = :frm"), {"to": db_to, "frm": db_from})
            updated = result.rowcount if result is not None else None

        flash(f'{updated} registros movidos de {visual_from} para {visual_to}.', 'success')
        return redirect(url_for('index'))
    except Exception as e:
        print(f"ERRO AO CORRIGIR DATA: {e}", flush=True)
        flash(f'Erro ao corrigir data: {str(e)}', 'error')
        return redirect(url_for('index'))


@app.route('/delete_date', methods=['POST'])
def delete_date():
    """Deletes all records for a given data_referencia. Use with caution."""
    target = request.form.get('target_date')
    confirm = request.form.get('confirm')
    if not target:
        flash('Data é obrigatória para exclusão.', 'error')
        return redirect(url_for('index'))
    if confirm != 'on':
        flash('Por segurança, marque a confirmação antes de excluir.', 'error')
        return redirect(url_for('index'))

    try:
        # Tenta interpretar com dayfirst=True (padrão do sistema)
        dt = pd.to_datetime(target, dayfirst=True)
        db_target = dt.strftime('%Y-%m-%d')
        visual = dt.strftime('%d/%m/%Y')
    except Exception as e:
        flash(f'Formato de data inválido: {str(e)}', 'error')
        return redirect(url_for('index'))

    try:
        with engine.begin() as conn:
            cnt = conn.execute(text("SELECT COUNT(*) FROM historico_ocupacao_completo WHERE data_referencia = :d"), {"d": db_target}).scalar()
            if cnt == 0:
                # Tenta interpretação alternativa (dia/mês trocados) para detectar divergências de entrada
                try:
                    dt_alt = pd.to_datetime(target, dayfirst=False)
                    alt_db = dt_alt.strftime('%Y-%m-%d')
                    alt_visual = dt_alt.strftime('%d/%m/%Y')
                    cnt_alt = conn.execute(text("SELECT COUNT(*) FROM historico_ocupacao_completo WHERE data_referencia = :d"), {"d": alt_db}).scalar()
                except Exception:
                    cnt_alt = 0

                if cnt_alt > 0:
                    # Exclui usando a interpretação alternativa e informa o usuário
                    res = conn.execute(text("DELETE FROM historico_ocupacao_completo WHERE data_referencia = :d"), {"d": alt_db})
                    deleted = res.rowcount if res is not None else None
                    flash(f'{deleted} registros excluídos para {alt_visual} (interpretado automaticamente).', 'success')
                    return redirect(url_for('index'))

                flash(f'Nenhum registro encontrado para {visual}.', 'error')
                return redirect(url_for('index'))

            res = conn.execute(text("DELETE FROM historico_ocupacao_completo WHERE data_referencia = :d"), {"d": db_target})
            deleted = res.rowcount if res is not None else None

        flash(f'{deleted} registros excluídos para {visual}.', 'success')
        return redirect(url_for('index'))
    except Exception as e:
        print(f"ERRO AO EXCLUIR DATA: {e}", flush=True)
        flash(f'Erro ao excluir registros: {str(e)}', 'error')
        return redirect(url_for('index'))

# ===== ROTAS API REST =====
@app.route('/api/version')
def api_version():
    """Retorna versão do sistema"""
    return {"version": get_version()}

@app.route('/api/stats')
def api_stats():
    """Retorna estatísticas gerais em JSON"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        selected_date = request.args.get('data')
        with engine.connect() as conn:
            # Busca histórico para pegar última data se necessário
            if not selected_date:
                sql_last = text("SELECT data_referencia FROM historico_ocupacao_completo ORDER BY data_referencia DESC LIMIT 1")
                result = conn.execute(sql_last).scalar()
                if result:
                    selected_date = result
                else:
                    return {"error": "Sem dados disponíveis"}, 404
            
            # Estatísticas
            sql_stats = text("""
                SELECT 
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'LIVRE' THEN 1 ELSE 0 END), 0) as vagos
                FROM historico_ocupacao_completo
                WHERE data_referencia = :data
            """)
            stats = conn.execute(sql_stats, {"data": selected_date}).mappings().fetchone()
            
            return {
                "total": int(stats['total']),
                "ocupados": int(stats['ocupados']),
                "impedidos": int(stats['impedidos']),
                "vagos": int(stats['vagos']),
                "data_referencia": str(selected_date)
            }
    except Exception as e:
        return {"error": str(e)}, 500

@app.route('/api/chart')
def api_chart():
    """Retorna dados para gráfico de evolução em JSON"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        selected_date = request.args.get('data')
        with engine.connect() as conn:
            if not selected_date:
                sql_last = text("SELECT data_referencia FROM historico_ocupacao_completo ORDER BY data_referencia DESC LIMIT 1")
                result = conn.execute(sql_last).scalar()
                if result:
                    selected_date = result
                else:
                    return {"error": "Sem dados disponíveis"}, 404
            
            sql_chart = text("""
                SELECT 
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados
                FROM historico_ocupacao_completo
                WHERE data_referencia <= :data
                GROUP BY data_referencia
                ORDER BY data_referencia DESC
                LIMIT 7
            """)
            chart_data = conn.execute(sql_chart, {"data": selected_date}).mappings().all()
            chart_data = list(reversed(chart_data))
            
            return {
                "labels": [row['dia'] for row in chart_data],
                "data": [int(row['ocupados']) for row in chart_data]
            }
    except Exception as e:
        return {"error": str(e)}, 500

@app.route('/api/history')
def api_history():
    """Retorna histórico de importações em JSON"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            sql_history = text("""
                SELECT 
                    data_referencia,
                    DATE_FORMAT(data_referencia, '%d/%m/%Y') as data_formatada,
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados
                FROM historico_ocupacao_completo
                GROUP BY data_referencia
                ORDER BY data_referencia DESC
            """)
            history_list = conn.execute(sql_history).mappings().all()
            
            return {
                "history": [
                    {
                        "data_referencia": str(row['data_referencia']),
                        "data_formatada": row['data_formatada'],
                        "total": int(row['total']),
                        "ocupados": int(row['ocupados'])
                    }
                    for row in history_list
                ]
            }
    except Exception as e:
        return {"error": str(e)}, 500

@app.route('/api/painel/stats')
def api_painel_stats():
    """Retorna estatísticas do painel de ocupação com suporte a filtros"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros da query string
        predio = request.args.get('predio')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')
        
        with engine.connect() as conn:
            # Monta condições WHERE dinamicamente
            where_conditions = []
            params = {}
            
            # Filtro de prédio (baseado no num_enf)
            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")
            
            # Filtro de período
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            elif not periodo_inicio and not periodo_fim and not mes:
                # Se não tem filtros de data, usa última data
                sql_last_date = text("SELECT MAX(data_referencia) as ultima_data FROM historico_ocupacao_completo")
                ultima_data = conn.execute(sql_last_date).scalar()
                if not ultima_data:
                    return {"error": "Sem dados disponíveis"}, 404
                where_conditions.append("data_referencia = :ultima_data")
                params['ultima_data'] = ultima_data
            
            # Filtro de mês
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            # Filtro de clínica
            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica
            
            # Monta SQL com WHERE dinâmico
            where_clause = " AND ".join(where_conditions) if where_conditions else "1=1"
            
            sql_stats = text(f"""
                SELECT 
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'LIVRE' THEN 1 ELSE 0 END), 0) as livres,
                    COALESCE(SUM(CASE WHEN status_leito = 'CEDIDO' THEN 1 ELSE 0 END), 0) as cedidos,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'RESERVADO' THEN 1 ELSE 0 END), 0) as reservados,
                    COUNT(*) as total
                FROM historico_ocupacao_completo
                WHERE {where_clause}
            """)
            stats = conn.execute(sql_stats, params).mappings().fetchone()
            
            return {
                "ocupados": int(stats['ocupados']),
                "livres": int(stats['livres']),
                "cedidos": int(stats['cedidos']),
                "impedidos": int(stats['impedidos']),
                "reservados": int(stats['reservados']),
                "total": int(stats['total'])
            }
    except Exception as e:
        return {"error": str(e)}, 500

@app.route('/api/painel/evolucao')
def api_painel_evolucao():
    """Retorna evolução mensal para gráfico com suporte a filtros"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros
        predio = request.args.get('predio')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')
        
        with engine.connect() as conn:
            # Monta condições WHERE
            where_conditions = []
            params = {}
            
            # Filtro de prédio
            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")
            
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica
            
            where_clause = " AND ".join(where_conditions) if where_conditions else "1=1"
            
            sql_evolucao = text(f"""
                SELECT DATE_FORMAT(data_referencia, '%Y-%m') as mes,
                    SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END) as ocupados,
                    COUNT(*) as total
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY mes
                ORDER BY mes
            """)
            evolucao = conn.execute(sql_evolucao, params).mappings().all()
            
            return {
                "labels": [row['mes'] for row in evolucao],
                "data": [round((int(row['ocupados']) / int(row['total'])) * 100, 1) for row in evolucao]
            }
    except Exception as e:
        return {"error": str(e)}, 500

@app.route('/api/painel/clinicas')
def api_painel_clinicas():
    """Retorna dados por clínica com suporte a filtros"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros
        predio = request.args.get('predio')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')
        
        with engine.connect() as conn:
            # Monta condições WHERE
            where_conditions = []
            params = {}
            
            # Filtro de prédio
            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")
            
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica
            
            where_clause = " AND ".join(where_conditions) if where_conditions else "1=1"
            
            sql_clinica = text(f"""
                SELECT nome_enfermaria,
                    SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END) as ocupados,
                    COUNT(*) as total
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY nome_enfermaria
                ORDER BY nome_enfermaria
            """)
            clinicas = conn.execute(sql_clinica, params).mappings().all()
            
            return {
                "labels": [row['nome_enfermaria'] for row in clinicas],
                "data": [round((int(row['ocupados']) / int(row['total'])) * 100, 1) for row in clinicas]
            }
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/painel/impedimentos')
def api_painel_impedimentos():
    """Retorna motivos de impedimento e contagens, respeitando filtros."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        predio = request.args.get('predio')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')

        with engine.connect() as conn:
            where_conditions = ["status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%'"]
            params = {}

            # Handle date filtering
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
                # If mes is also provided with period, add month filter too
                if mes:
                    year, month = _parse_mes_param(conn, mes)
                    if year is None or month is None:
                        return {"error": "Filtro de mes invalido"}, 400
                    where_conditions.append("MONTH(data_referencia) = :mes")
                    where_conditions.append("YEAR(data_referencia) = :ano")
                    params['mes'] = month
                    params['ano'] = year
            elif mes:
                # Month-only filter: use same year as latest data
                sql_last = text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")
                last = conn.execute(sql_last).scalar()
                if not last:
                    return {"error": "Sem dados disponíveis"}, 404
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                params['mes'] = month
                params['ano'] = year
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
            else:
                # No filters: default to last 14 days
                sql_last = text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")
                last = conn.execute(sql_last).scalar()
                if not last:
                    return {"error": "Sem dados disponíveis"}, 404
                where_conditions.append("data_referencia BETWEEN DATE_SUB(:last, INTERVAL 13 DAY) AND :last")
                params['last'] = last

            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica

            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")

            where_clause = " AND ".join(where_conditions) if where_conditions else "1=1"

            sql = text(f"""
                SELECT COALESCE(NULLIF(TRIM(motivo_impedimento), ''), 'Sem motivo informado') as motivo,
                       COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY motivo
                ORDER BY cnt DESC
                LIMIT 20
            """)

            rows = conn.execute(sql, params).mappings().all()

            return {"labels": [r['motivo'] for r in rows], "data": [int(r['cnt']) for r in rows]}

    except Exception as e:
        return {"error": str(e)}, 500


# ========== EMERGÊNCIA ENDPOINTS ==========

@app.route('/api/emergencia/stats')
def api_emergencia_stats():
    """Retorna estatísticas do painel de emergência com suporte a filtros"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros da query string
        enfermaria = request.args.get('enfermaria')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        
        with engine.connect() as conn:
            # Monta condições WHERE dinamicamente
            where_conditions = []
            params = {}
            
            # FILTRO PRINCIPAL: Apenas enfermarias de emergência
            ward_list = "', '".join(EMERGENCY_WARDS)
            where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
            
            # Filtro de enfermaria específica
            if enfermaria and enfermaria in EMERGENCY_WARDS:
                where_conditions = [f"nome_enfermaria = :enfermaria"]
                params['enfermaria'] = enfermaria
            
            # Filtro de período
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            elif not periodo_inicio and not periodo_fim and not mes:
                # Se não tem filtros de data, usa última data disponível
                sql_last_date = text("SELECT MAX(data_referencia) as ultima_data FROM historico_ocupacao_completo")
                ultima_data = conn.execute(sql_last_date).scalar()
                if not ultima_data:
                    return {"error": "Sem dados disponíveis"}, 404
                where_conditions.append("data_referencia = :ultima_data")
                params['ultima_data'] = ultima_data
            
            # Filtro de mês
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            # Monta SQL com WHERE dinâmico
            where_clause = " AND " + " AND ".join(where_conditions) if where_conditions else ""
            
            # Estatísticas gerais
            sql_stats = text(f"""
                SELECT 
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'LIVRE' THEN 1 ELSE 0 END), 0) as livres,
                    COALESCE(SUM(CASE WHEN status_leito = 'CEDIDO' THEN 1 ELSE 0 END), 0) as cedidos,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'RESERVADO' THEN 1 ELSE 0 END), 0) as reservados,
                    COUNT(*) as total
                FROM historico_ocupacao_completo
                WHERE 1=1 {where_clause}
            """)
            stats = conn.execute(sql_stats, params).mappings().fetchone()
            
            ocupados = int(stats['ocupados'])
            total = int(stats['total'])
            taxa_ocupacao = round((ocupados / total * 100), 2) if total > 0 else 0
            
            # Estatística pediátrica
            sql_ped = text(f"""
                SELECT 
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ped_ocupados,
                    COUNT(*) as ped_total
                FROM historico_ocupacao_completo
                WHERE nome_enfermaria = 'CLINICA REFERENCIADA - PED' {where_clause.replace(f"nome_enfermaria IN ('{ward_list}')", '1=1')}
            """)
            ped_stats = conn.execute(sql_ped, params).mappings().fetchone()
            ped_ocupados = int(ped_stats['ped_ocupados'])
            ped_total = int(ped_stats['ped_total'])
            
            # Tempo médio de permanência (apenas ocupados)
            sql_tempo = text(f"""
                SELECT 
                    AVG(TIMESTAMPDIFF(DAY, data_internacao, data_referencia)) as tempo_medio
                FROM historico_ocupacao_completo
                WHERE status_leito = 'OCUPADO' 
                    AND data_internacao IS NOT NULL
                    AND nome_enfermaria IN ('{ward_list}') {where_clause.replace(f"nome_enfermaria IN ('{ward_list}')", '1=1')}
            """)
            tempo_result = conn.execute(sql_tempo, params).scalar()
            tempo_medio = round(float(tempo_result), 1) if tempo_result else 0
            
            # Rotatividade (pacientes distintos / total leitos nos últimos 30 dias)
            sql_rotatividade = text(f"""
                SELECT 
                    COUNT(DISTINCT nome_paciente) as pacientes_distintos
                FROM historico_ocupacao_completo
                WHERE status_leito = 'OCUPADO'
                    AND nome_paciente IS NOT NULL
                    AND nome_paciente != ''
                    AND data_referencia >= DATE_SUB((SELECT MAX(data_referencia) FROM historico_ocupacao_completo), INTERVAL 30 DAY)
                    AND nome_enfermaria IN ('{ward_list}')
            """)
            pacientes_distintos = conn.execute(sql_rotatividade).scalar() or 0
            rotatividade = round((pacientes_distintos / total), 2) if total > 0 else 0
            
            return {
                "ocupados": ocupados,
                "livres": int(stats['livres']),
                "cedidos": int(stats['cedidos']),
                "impedidos": int(stats['impedidos']),
                "reservados": int(stats['reservados']),
                "total": total,
                "taxa_ocupacao": taxa_ocupacao,
                "pediatrico_ocupados": ped_ocupados,
                "pediatrico_total": ped_total,
                "tempo_medio_permanencia": tempo_medio,
                "rotatividade": rotatividade
            }
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/evolucao')
def api_emergencia_evolucao():
    """Retorna evolução temporal da taxa de ocupação para emergência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros
        enfermaria = request.args.get('enfermaria')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        
        with engine.connect() as conn:
            # Monta condições WHERE
            where_conditions = []
            params = {}
            
            # FILTRO PRINCIPAL: Apenas enfermarias de emergência
            ward_list = "', '".join(EMERGENCY_WARDS)
            where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
            
            if enfermaria and enfermaria in EMERGENCY_WARDS:
                where_conditions = [f"nome_enfermaria = :enfermaria"]
                params['enfermaria'] = enfermaria
            
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            elif not mes:
                # Padrão: últimos 30 dias
                where_conditions.append("data_referencia >= DATE_SUB((SELECT MAX(data_referencia) FROM historico_ocupacao_completo), INTERVAL 30 DAY)")
            
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            where_clause = " AND " + " AND ".join(where_conditions) if where_conditions else ""
            
            sql_evolucao = text(f"""
                SELECT 
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END) as ocupados,
                    COUNT(*) as total
                FROM historico_ocupacao_completo
                WHERE 1=1 {where_clause}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql_evolucao, params).mappings().all()
            
            labels = [r['dia'] for r in rows]
            data = [round((r['ocupados'] / r['total'] * 100), 2) if r['total'] > 0 else 0 for r in rows]
            
            return {"labels": labels, "data": data}
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/enfermarias')
def api_emergencia_enfermarias():
    """Retorna ocupação por enfermaria de emergência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        
        with engine.connect() as conn:
            # Monta condições WHERE
            where_conditions = []
            params = {}
            
            # FILTRO PRINCIPAL: Apenas enfermarias de emergência
            ward_list = "', '".join(EMERGENCY_WARDS)
            where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
            
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            elif not mes:
                # Padrão: últimos 30 dias
                where_conditions.append("data_referencia >= DATE_SUB((SELECT MAX(data_referencia) FROM historico_ocupacao_completo), INTERVAL 30 DAY)")
            
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            where_clause = " AND " + " AND ".join(where_conditions) if where_conditions else ""
            
            sql_enfermarias = text(f"""
                SELECT 
                    nome_enfermaria,
                    SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END) as ocupados,
                    COUNT(*) as total
                FROM historico_ocupacao_completo
                WHERE 1=1 {where_clause}
                GROUP BY nome_enfermaria
                ORDER BY nome_enfermaria
            """)
            rows = conn.execute(sql_enfermarias, params).mappings().all()
            
            labels = [r['nome_enfermaria'] for r in rows]
            ocupados = [int(r['ocupados']) for r in rows]
            totals = [int(r['total']) for r in rows]
            
            return {
                "labels": labels,
                "ocupados": ocupados,
                "total": totals
            }
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/pacientes')
def api_emergencia_pacientes():
    """Retorna lista paginada de pacientes internados na emergência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Paginação
        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 20))
        offset = (page - 1) * per_page
        
        # Filtros
        enfermaria = request.args.get('enfermaria')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        
        with engine.connect() as conn:
            # Monta condições WHERE
            where_conditions = []
            params = {}
            
            # FILTRO PRINCIPAL: Apenas enfermarias de emergência e ocupados
            ward_list = "', '".join(EMERGENCY_WARDS)
            where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
            where_conditions.append("status_leito = 'OCUPADO'")
            
            if enfermaria and enfermaria in EMERGENCY_WARDS:
                where_conditions.append("nome_enfermaria = :enfermaria")
                params['enfermaria'] = enfermaria
            
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            elif mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            else:
                # Padrão: última data disponível
                sql_last_date = text("SELECT MAX(data_referencia) as ultima_data FROM historico_ocupacao_completo")
                ultima_data = conn.execute(sql_last_date).scalar()
                if ultima_data:
                    where_conditions.append("data_referencia = :ultima_data")
                    params['ultima_data'] = ultima_data
            
            where_clause = " AND " + " AND ".join(where_conditions) if where_conditions else ""
            
            # Count total
            sql_count = text(f"""
                SELECT COUNT(DISTINCT CONCAT(nome_paciente, '|', nome_enfermaria)) as total
                FROM historico_ocupacao_completo
                WHERE 1=1 {where_clause}
            """)
            total_count = conn.execute(sql_count, params).scalar() or 0
            
            # Get paginated data
            params['offset'] = offset
            params['per_page'] = per_page
            
            sql_pacientes = text(f"""
                SELECT 
                    nome_paciente,
                    sexo,
                    idade,
                    nome_enfermaria,
                    data_internacao,
                    TIMESTAMPDIFF(DAY, data_internacao, data_referencia) as dias_permanencia
                FROM historico_ocupacao_completo
                WHERE 1=1 {where_clause}
                ORDER BY dias_permanencia DESC
                LIMIT :per_page OFFSET :offset
            """)
            rows = conn.execute(sql_pacientes, params).mappings().all()
            
            pacientes = []
            for r in rows:
                pacientes.append({
                    "nome": r['nome_paciente'] or 'Não informado',
                    "sexo": r['sexo'] or '-',
                    "idade": int(r['idade']) if r['idade'] else 0,
                    "enfermaria": r['nome_enfermaria'],
                    "data_internacao": r['data_internacao'].strftime('%d/%m/%Y') if r['data_internacao'] else '-',
                    "dias_permanencia": int(r['dias_permanencia']) if r['dias_permanencia'] else 0
                })
            
            return {
                "pacientes": pacientes,
                "total": total_count,
                "page": page,
                "per_page": per_page,
                "total_pages": (total_count + per_page - 1) // per_page
            }
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/export')
def api_emergencia_export():
    """Exporta dados dos pacientes de emergência para Excel"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        # Captura filtros
        enfermaria = request.args.get('enfermaria')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        
        with engine.connect() as conn:
            # Monta condições WHERE
            where_conditions = []
            params = {}
            
            # FILTRO PRINCIPAL: Apenas enfermarias de emergência
            ward_list = "', '".join(EMERGENCY_WARDS)
            where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
            
            # Filtro de enfermaria específica
            if enfermaria and enfermaria in EMERGENCY_WARDS:
                where_conditions = [f"nome_enfermaria = :enfermaria"]
                params['enfermaria'] = enfermaria
            
            # Filtro de período
            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            elif not periodo_inicio and not periodo_fim and not mes:
                # Se não tem filtros de data, usa última data
                sql_last_date = text("SELECT MAX(data_referencia) as ultima_data FROM historico_ocupacao_completo")
                ultima_data = conn.execute(sql_last_date).scalar()
                if not ultima_data:
                    return {"error": "Sem dados disponíveis"}, 404
                where_conditions.append("data_referencia = :ultima_data")
                params['ultima_data'] = ultima_data
            
            # Filtro de mês
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            
            where_clause = " AND " + " AND ".join(where_conditions) if where_conditions else ""
            
            # Query para buscar pacientes ocupados
            sql_export = text(f"""
                SELECT 
                    nome_paciente as paciente,
                    prontuario,
                    sexo,
                    idade,
                    nome_enfermaria as enfermaria,
                    DATE_FORMAT(data_internacao, '%d/%m/%Y') as data_internacao,
                    TIMESTAMPDIFF(DAY, data_internacao, data_referencia) as dias_permanencia,
                    status_leito,
                    DATE_FORMAT(data_referencia, '%d/%m/%Y') as data_referencia
                FROM historico_ocupacao_completo
                WHERE status_leito = 'OCUPADO'
                    AND nome_paciente IS NOT NULL
                    AND nome_paciente != ''
                    {where_clause}
                ORDER BY nome_enfermaria, dias_permanencia DESC
            """)
            
            rows = conn.execute(sql_export, params).mappings().all()
            
            if not rows:
                return {"error": "Sem dados para exportar"}, 404
            
            # Converte para DataFrame
            df = pd.DataFrame([{
                'Paciente': r['paciente'],
                'Prontuário': r['prontuario'],
                'Sexo': r['sexo'],
                'Idade': int(r['idade']) if r['idade'] is not None else None,
                'Enfermaria': r['enfermaria'],
                'Data Internação': r['data_internacao'],
                'Dias Permanência': int(r['dias_permanencia']) if r['dias_permanencia'] is not None else 0,
                'Status Leito': r['status_leito'],
                'Data Referência': r['data_referencia']
            } for r in rows])
            
            # Cria arquivo Excel
            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='Emergência')
            output.seek(0)
            
            # Nome do arquivo com data
            filename = f"emergencia_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            return send_file(
                output,
                download_name=filename,
                as_attachment=True,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
    except Exception as e:
        return {"error": str(e)}, 500


# Helper function for emergency range context
def _get_emergencia_range_context(conn, args):
    """Helper para construir contexto de filtros para emergência com suporte a período"""
    enfermaria = args.get('enfermaria')
    periodo_inicio = args.get('periodo_inicio')
    periodo_fim = args.get('periodo_fim')
    mes = args.get('mes')
    
    where_conditions = []
    params = {}
    
    # FILTRO PRINCIPAL: Apenas enfermarias de emergência
    ward_list = "', '".join(EMERGENCY_WARDS)
    where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
    
    # Filtro de enfermaria específica
    if enfermaria and enfermaria in EMERGENCY_WARDS:
        where_conditions = [f"nome_enfermaria = :enfermaria"]
        params['enfermaria'] = enfermaria
    
    # Filtro de período
    if periodo_inicio and periodo_fim:
        where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
        params['periodo_inicio'] = periodo_inicio
        params['periodo_fim'] = periodo_fim
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is None or month is None:
            return None
        where_conditions.append("MONTH(data_referencia) = :mes")
        where_conditions.append("YEAR(data_referencia) = :ano")
        params['mes'] = month
        params['ano'] = year
    else:
        # Padrão: últimos 14 dias
        max_date = conn.execute(text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")).scalar()
        if not max_date:
            return None
        where_conditions.append("data_referencia BETWEEN DATE_SUB(:max_date, INTERVAL 13 DAY) AND :max_date")
        params['max_date'] = max_date
    
    return {
        "where": " AND ".join(where_conditions) if where_conditions else "1=1",
        "params": params
    }


def _get_emergencia_profile_snapshot_context(conn, args):
    enfermaria = args.get('enfermaria')
    periodo_inicio = args.get('periodo_inicio')
    periodo_fim = args.get('periodo_fim')
    mes = args.get('mes')

    ward_list = "', '".join(EMERGENCY_WARDS)
    base_conditions = [f"nome_enfermaria IN ('{ward_list}')"]
    base_params = {}

    if enfermaria and enfermaria in EMERGENCY_WARDS:
        base_conditions = ["nome_enfermaria = :enfermaria"]
        base_params['enfermaria'] = enfermaria

    selected_date = None

    if periodo_fim:
        selected_date = periodo_fim
    elif periodo_inicio:
        selected_date = periodo_inicio
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is not None and month is not None:
            month_conditions = list(base_conditions)
            month_conditions.append("MONTH(data_referencia) = :mes")
            month_conditions.append("YEAR(data_referencia) = :ano")
            month_params = dict(base_params)
            month_params['mes'] = month
            month_params['ano'] = year
            month_where = " AND ".join(month_conditions) if month_conditions else "1=1"
            sql_month_last = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {month_where}")
            selected_date = conn.execute(sql_month_last, month_params).scalar()

    if not selected_date:
        latest_where = " AND ".join(base_conditions) if base_conditions else "1=1"
        sql_last = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {latest_where}")
        selected_date = conn.execute(sql_last, base_params).scalar()

    if not selected_date:
        return None

    snapshot_conditions = list(base_conditions)
    snapshot_conditions.append("data_referencia = :data_referencia")
    snapshot_conditions.append("status_leito = 'OCUPADO'")
    snapshot_params = dict(base_params)
    snapshot_params['data_referencia'] = selected_date

    return {
        "selected_date": selected_date,
        "where": " AND ".join(snapshot_conditions),
        "params": snapshot_params
    }


def _get_emergencia_profile_range_context(conn, args):
    enfermaria = args.get('enfermaria')
    periodo_inicio = args.get('periodo_inicio')
    periodo_fim = args.get('periodo_fim')
    mes = args.get('mes')

    ward_list = "', '".join(EMERGENCY_WARDS)
    where_conditions = ["status_leito = 'OCUPADO'"]
    params = {}

    if enfermaria and enfermaria in EMERGENCY_WARDS:
        where_conditions.append("nome_enfermaria = :enfermaria")
        params['enfermaria'] = enfermaria
    else:
        where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")

    if periodo_inicio and periodo_fim:
        where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
        params['periodo_inicio'] = periodo_inicio
        params['periodo_fim'] = periodo_fim
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is None or month is None:
            return None
        where_conditions.append("MONTH(data_referencia) = :mes")
        where_conditions.append("YEAR(data_referencia) = :ano")
        params['mes'] = month
        params['ano'] = year
    else:
        base_where = " AND ".join(where_conditions) if where_conditions else "1=1"
        sql_max = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {base_where}")
        max_date = conn.execute(sql_max, params).scalar()
        if not max_date:
            return None
        where_conditions.append("data_referencia BETWEEN DATE_SUB(:max_date, INTERVAL 13 DAY) AND :max_date")
        params['max_date'] = max_date

    return {
        "where": " AND ".join(where_conditions) if where_conditions else "1=1",
        "params": params
    }


@app.route('/api/emergencia/perfil-resumo')
def api_emergencia_perfil_resumo():
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            context = _get_emergencia_profile_snapshot_context(conn, request.args)
            if not context:
                return {"error": "Sem dados disponíveis"}, 404

            params = dict(context['params'])
            params['ref_date'] = context['selected_date']
            sql = text(_patient_profile_query(context['where']))
            rows = conn.execute(sql, params).mappings().all()

            total = len(rows)
            if total == 0:
                return jsonify({
                    "data_referencia": str(context['selected_date']),
                    "total_pacientes": 0,
                    "sexo_m_pct": 0,
                    "sexo_f_pct": 0,
                    "idade_media": 0,
                    "idade_mediana": 0,
                    "tempo_medio": 0,
                    "tempo_mediano": 0,
                    "longa_permanencia_pct": 0
                })

            idades = [int(r['idade']) for r in rows if r['idade'] is not None]
            tempos = [int(r['dias']) if r['dias'] is not None else 0 for r in rows]

            sex_norm = []
            for r in rows:
                sexo_raw = (r['sexo'] or '').strip().upper()
                if sexo_raw in ('M', 'MASCULINO'):
                    sex_norm.append('M')
                elif sexo_raw in ('F', 'FEMININO'):
                    sex_norm.append('F')
                else:
                    sex_norm.append('NI')

            def median(values):
                if not values:
                    return 0
                ordered = sorted(values)
                n = len(ordered)
                m = n // 2
                if n % 2 == 0:
                    return round((ordered[m - 1] + ordered[m]) / 2, 1)
                return ordered[m]

            male_count = sum(1 for s in sex_norm if s == 'M')
            female_count = sum(1 for s in sex_norm if s == 'F')
            longa_count = sum(1 for d in tempos if d > 30)

            return jsonify({
                "data_referencia": str(context['selected_date']),
                "total_pacientes": total,
                "sexo_m_pct": round((male_count / total) * 100, 1),
                "sexo_f_pct": round((female_count / total) * 100, 1),
                "idade_media": round(sum(idades) / len(idades), 1) if idades else 0,
                "idade_mediana": median(idades),
                "tempo_medio": round(sum(tempos) / len(tempos), 1) if tempos else 0,
                "tempo_mediano": median(tempos),
                "longa_permanencia_pct": round((longa_count / total) * 100, 1)
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/perfil-graficos')
def api_emergencia_perfil_graficos():
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            snapshot_context = _get_emergencia_profile_snapshot_context(conn, request.args)
            if not snapshot_context:
                return {"error": "Sem dados disponíveis"}, 404

            snapshot_params = dict(snapshot_context['params'])
            snapshot_params['ref_date'] = snapshot_context['selected_date']
            sql_profile = text(_patient_profile_query(snapshot_context['where']))
            rows = conn.execute(sql_profile, snapshot_params).mappings().all()

            range_context = _get_emergencia_profile_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_series = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(DISTINCT COALESCE(
                        NULLIF(TRIM(cns_paciente), ''),
                        NULLIF(TRIM(prontuario), ''),
                        NULLIF(TRIM(aih_paciente), ''),
                        NULLIF(TRIM(nome_paciente), ''),
                        CONCAT('LEITO-', num_enf, '-', leito)
                    )) as pacientes_ativos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            series_rows = conn.execute(sql_series, range_context['params']).mappings().all()

            sexo_counts = {'Masculino': 0, 'Feminino': 0, 'Não informado': 0}
            faixa_counts = {'0-17': 0, '18-39': 0, '40-59': 0, '60-79': 0, '80+': 0}
            hist_counts = {'0-7': 0, '8-14': 0, '15-30': 0, '31-60': 0, '61-90': 0, '>90': 0}
            enfermaria_counts = {}
            enfermaria_los = {}

            for row in rows:
                sexo_raw = (row['sexo'] or '').strip().upper()
                if sexo_raw in ('M', 'MASCULINO'):
                    sexo_counts['Masculino'] += 1
                elif sexo_raw in ('F', 'FEMININO'):
                    sexo_counts['Feminino'] += 1
                else:
                    sexo_counts['Não informado'] += 1

                idade = int(row['idade']) if row['idade'] is not None else None
                if idade is not None:
                    if idade <= 17:
                        faixa_counts['0-17'] += 1
                    elif idade <= 39:
                        faixa_counts['18-39'] += 1
                    elif idade <= 59:
                        faixa_counts['40-59'] += 1
                    elif idade <= 79:
                        faixa_counts['60-79'] += 1
                    else:
                        faixa_counts['80+'] += 1

                dias = int(row['dias']) if row['dias'] is not None else 0
                if dias <= 7:
                    hist_counts['0-7'] += 1
                elif dias <= 14:
                    hist_counts['8-14'] += 1
                elif dias <= 30:
                    hist_counts['15-30'] += 1
                elif dias <= 60:
                    hist_counts['31-60'] += 1
                elif dias <= 90:
                    hist_counts['61-90'] += 1
                else:
                    hist_counts['>90'] += 1

                enfermaria = row['nome_enfermaria'] or 'Sem enfermaria'
                enfermaria_counts[enfermaria] = enfermaria_counts.get(enfermaria, 0) + 1
                if enfermaria not in enfermaria_los:
                    enfermaria_los[enfermaria] = {'sum': 0, 'count': 0}
                enfermaria_los[enfermaria]['sum'] += dias
                enfermaria_los[enfermaria]['count'] += 1

            top_enfermarias = sorted(enfermaria_counts.items(), key=lambda item: item[1], reverse=True)[:10]
            enfermaria_avg_los = []
            for enfermaria, stats in enfermaria_los.items():
                avg_days = round(stats['sum'] / stats['count'], 1) if stats['count'] > 0 else 0
                enfermaria_avg_los.append((enfermaria, avg_days))
            top_enfermarias_los = sorted(enfermaria_avg_los, key=lambda item: item[1], reverse=True)[:10]

            return jsonify({
                "serie_pacientes": {
                    "labels": [r['dia'] for r in series_rows],
                    "data": [int(r['pacientes_ativos']) for r in series_rows]
                },
                "sexo": {
                    "labels": list(sexo_counts.keys()),
                    "data": list(sexo_counts.values())
                },
                "faixa_etaria": {
                    "labels": list(faixa_counts.keys()),
                    "data": list(faixa_counts.values())
                },
                "hist_permanencia": {
                    "labels": list(hist_counts.keys()),
                    "data": list(hist_counts.values())
                },
                "top_enfermarias": {
                    "labels": [c[0] for c in top_enfermarias],
                    "data": [c[1] for c in top_enfermarias]
                },
                "permanencia_enfermaria": {
                    "labels": [c[0] for c in top_enfermarias_los],
                    "data": [c[1] for c in top_enfermarias_los]
                }
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/kpis-complementares')
def api_emergencia_kpis_complementares():
    """Retorna KPIs complementares do perfil da emergência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql_stats = text(f"""
                SELECT
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' AND UPPER(TRIM(cronico)) IN ('SIM','S','1','TRUE','T','YES') THEN 1 ELSE 0 END), 0) as cronicos,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' AND NULLIF(TRIM(modo_ventilatorio), '') IS NOT NULL THEN 1 ELSE 0 END), 0) as ventilacao,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' AND UPPER(TRIM(inserido_no_trs)) IN ('SIM','S','1','TRUE','T','YES') THEN 1 ELSE 0 END), 0) as trs
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
            """)
            stats = conn.execute(sql_stats, range_context['params']).mappings().fetchone()
            
            sql_top_enf = text(f"""
                SELECT COALESCE(NULLIF(TRIM(nome_enfermaria), ''), 'Sem clínica') as enfermaria,
                       COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
                GROUP BY enfermaria
                ORDER BY cnt DESC
                LIMIT 5
            """)
            top_rows = conn.execute(sql_top_enf, range_context['params']).mappings().all()
            
            total = int(stats['total'] or 0)
            ocupados = int(stats['ocupados'] or 0)
            impedidos = int(stats['impedidos'] or 0)
            cronicos = int(stats['cronicos'] or 0)
            ventilacao = int(stats['ventilacao'] or 0)
            trs = int(stats['trs'] or 0)
            
            return jsonify({
                "impedidos_pct": round((impedidos / total) * 100, 1) if total > 0 else 0,
                "top_enfermarias": [{"label": r['enfermaria'], "count": int(r['cnt'])} for r in top_rows],
                "cronicos_pct": round((cronicos / ocupados) * 100, 1) if ocupados > 0 else 0,
                "ventilacao_pct": round((ventilacao / ocupados) * 100, 1) if ocupados > 0 else 0,
                "trs_pct": round((trs / ocupados) * 100, 1) if ocupados > 0 else 0
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/impedimentos-serie')
def api_emergencia_impedimentos_serie():
    """Retorna série temporal da taxa de impedimentos por dia (%)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql_impedimentos = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' THEN 1 ELSE 0 END), 0) as impedidos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql_impedimentos, range_context['params']).mappings().all()
            
            return jsonify({
                "labels": [r['dia'] for r in rows],
                "data": [round((int(r['impedidos']) / int(r['total'])) * 100, 1) if int(r['total']) > 0 else 0 for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/status-serie')
def api_emergencia_status_serie():
    """Retorna série temporal de status de leitos (ocupado/livre/impedido/etc)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql_status = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COALESCE(SUM(CASE WHEN status_leito = 'LIVRE' THEN 1 ELSE 0 END), 0) as livres,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'CEDIDO' THEN 1 ELSE 0 END), 0) as cedidos,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'RESERVADO' THEN 1 ELSE 0 END), 0) as reservados
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql_status, range_context['params']).mappings().all()
            
            return jsonify({
                "labels": [r['dia'] for r in rows],
                "datasets": [
                    {"label": "Livres", "data": [int(r['livres']) for r in rows]},
                    {"label": "Ocupados", "data": [int(r['ocupados']) for r in rows]},
                    {"label": "Cedido", "data": [int(r['cedidos']) for r in rows]},
                    {"label": "Impedidos", "data": [int(r['impedidos']) for r in rows]},
                    {"label": "Reservados", "data": [int(r['reservados']) for r in rows]}
                ]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/impedimentos-top')
def api_emergencia_impedimentos_top():
    """Retorna top motivos de impedimento na emergência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql_top = text(f"""
                SELECT
                    COALESCE(NULLIF(TRIM(motivo_impedimento), ''), 'Não informado') as motivo,
                    COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                      AND (status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%')
                GROUP BY motivo
                ORDER BY cnt DESC
                LIMIT 10
            """)
            rows = conn.execute(sql_top, range_context['params']).mappings().all()
            
            return jsonify({
                "labels": [r['motivo'] for r in rows],
                "data": [int(r['cnt']) for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/ocupacao-heatmap')
def api_emergencia_ocupacao_heatmap():
    """Retorna heatmap de ocupação: dia da semana × enfermaria"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql_heatmap = text(f"""
                SELECT
                    DAYOFWEEK(data_referencia) as dia_num,
                    COALESCE(NULLIF(TRIM(nome_enfermaria), ''), 'Sem clínica') as enfermaria,
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY dia_num, enfermaria
                ORDER BY dia_num, enfermaria
            """)
            rows = conn.execute(sql_heatmap, range_context['params']).mappings().all()
            
            # Mapear dias da semana
            dias_map = {1: 'Dom', 2: 'Seg', 3: 'Ter', 4: 'Qua', 5: 'Qui', 6: 'Sex', 7: 'Sáb'}
            
            # Estruturar dados para heatmap
            enfermarias_set = set()
            for r in rows:
                enfermarias_set.add(r['enfermaria'])
            
            enfermarias = sorted(list(enfermarias_set))
            dias = ['Dom', 'Seg', 'Ter', 'Qua', 'Qui', 'Sex', 'Sáb']
            
            # Criar matriz de ocupação
            heatmap_data = []
            for enf in enfermarias:
                enf_data = {'enfermaria': enf, 'valores': []}
                for dia_idx in range(1, 8):
                    matching = [r for r in rows if r['enfermaria'] == enf and r['dia_num'] == dia_idx]
                    if matching:
                        taxa = round((int(matching[0]['ocupados']) / int(matching[0]['total'])) * 100, 1) if int(matching[0]['total']) > 0 else 0
                    else:
                        taxa = 0
                    enf_data['valores'].append(taxa)
                heatmap_data.append(enf_data)
            
            return jsonify({
                "dias": dias,
                "enfermarias": enfermarias,
                "data": heatmap_data
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/kpis-avancados')
def api_emergencia_kpis_avancados():
    """Retorna KPIs avançados da emergência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql_rot = text(f"""
                SELECT
                    COUNT(DISTINCT COALESCE(
                        NULLIF(TRIM(cns_paciente), ''),
                        NULLIF(TRIM(prontuario), ''),
                        NULLIF(TRIM(aih_paciente), ''),
                        NULLIF(TRIM(nome_paciente), ''),
                        CONCAT('LEITO-', num_enf, '-', leito)
                    )) as pacientes,
                    COUNT(DISTINCT CONCAT(num_enf, '-', leito)) as leitos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
            """)
            rot = conn.execute(sql_rot, range_context['params']).mappings().fetchone()
            
            sql_motivo = text(f"""
                SELECT
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN NULLIF(TRIM(situacao_motivo_permanencia), '') IS NOT NULL THEN 1 ELSE 0 END), 0) as com_motivo
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
            """)
            motivo = conn.execute(sql_motivo, range_context['params']).mappings().fetchone()
            
            sql_reserva = text(f"""
                SELECT
                    AVG(TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito)) as media_dias
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                  AND data_sol_reserva IS NOT NULL
                  AND data_internacao_leito IS NOT NULL
                  AND TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito) >= 0
            """)
            reserva = conn.execute(sql_reserva, range_context['params']).scalar()
            
            pacientes = int(rot['pacientes'] or 0)
            leitos = int(rot['leitos'] or 0)
            total = int(motivo['total'] or 0)
            com_motivo = int(motivo['com_motivo'] or 0)
            
            return jsonify({
                "rotatividade": round((pacientes / leitos), 2) if leitos > 0 else 0,
                "motivo_permanencia_pct": round((com_motivo / total) * 100, 1) if total > 0 else 0,
                "tempo_medio_reserva": round(float(reserva), 1) if reserva is not None else None
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/rotatividade-serie')
def api_emergencia_rotatividade_serie():
    """Retorna série temporal de rotatividade por leito (pacientes / leito)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(DISTINCT COALESCE(
                        NULLIF(TRIM(cns_paciente), ''),
                        NULLIF(TRIM(prontuario), ''),
                        NULLIF(TRIM(aih_paciente), ''),
                        NULLIF(TRIM(nome_paciente), ''),
                        CONCAT('LEITO-', num_enf, '-', leito)
                    )) as pacientes,
                    COUNT(DISTINCT CONCAT(num_enf, '-', leito)) as leitos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql, range_context['params']).mappings().all()
            
            return jsonify({
                "labels": [r['dia'] for r in rows],
                "data": [round((int(r['pacientes']) / int(r['leitos'])), 2) if int(r['leitos']) > 0 else 0 for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/reserva-serie')
def api_emergencia_reserva_serie():
    """Retorna série temporal de tempo médio de reserva e taxa de ocupação"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            range_context = _get_emergencia_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404
            
            sql = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    AVG(CASE 
                        WHEN data_sol_reserva IS NOT NULL 
                             AND data_internacao_leito IS NOT NULL 
                             AND TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito) >= 0
                        THEN TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito)
                        ELSE NULL 
                    END) as tempo_reserva
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql, range_context['params']).mappings().all()
            
            return jsonify({
                "labels": [r['dia'] for r in rows],
                "taxa_ocupacao": [round((int(r['ocupados']) / int(r['total'])) * 100, 1) if int(r['total']) > 0 else 0 for r in rows],
                "tempo_reserva": [round(float(r['tempo_reserva']), 1) if r['tempo_reserva'] is not None else 0 for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/emergencia/longa-permanencia-ranking')
def api_emergencia_longa_permanencia_ranking():
    """Retorna ranking de longa permanência por enfermaria de emergência (>30 dias)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    
    try:
        with engine.connect() as conn:
            # Para este endpoint, usar última data
            enfermaria = request.args.get('enfermaria')
            
            where_conditions = []
            params = {}
            
            # FILTRO PRINCIPAL: Apenas enfermarias de emergência
            ward_list = "', '".join(EMERGENCY_WARDS)
            where_conditions.append(f"nome_enfermaria IN ('{ward_list}')")
            
            if enfermaria and enfermaria in EMERGENCY_WARDS:
                where_conditions = [f"nome_enfermaria = :enfermaria"]
                params['enfermaria'] = enfermaria
            
            # Pegar última data
            sql_last_date = text("SELECT MAX(data_referencia) as ultima_data FROM historico_ocupacao_completo")
            ultima_data = conn.execute(sql_last_date).scalar()
            if not ultima_data:
                return {"error": "Sem dados disponíveis"}, 404
            
            where_conditions.append("data_referencia = :ultima_data")
            params['ultima_data'] = ultima_data
            
            where_clause = " AND " + " AND ".join(where_conditions)
            
            sql = text(f"""
                SELECT 
                    nome_enfermaria as clinica,
                    COUNT(*) as cnt,
                    AVG(TIMESTAMPDIFF(DAY, data_internacao, data_referencia)) as media
                FROM historico_ocupacao_completo
                WHERE status_leito = 'OCUPADO'
                    AND data_internacao IS NOT NULL
                    AND TIMESTAMPDIFF(DAY, data_internacao, data_referencia) > 30
                    {where_clause}
                GROUP BY clinica
                ORDER BY cnt DESC
                LIMIT 10
            """)
            rows = conn.execute(sql, params).mappings().all()
            
            return jsonify({
                "labels": [r['clinica'] for r in rows],
                "contagem": [int(r['cnt']) for r in rows],
                "media_dias": [round(float(r['media']), 1) for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/disponibilidade')
def disponibilidade():
    return render_template('disponibilidade.html')


@app.route('/perfil_paciente/sexo')
def perfil_paciente_sexo():
    """Renderiza a subpágina de Perfil do Paciente — Sexo"""
    return render_template('perfil_paciente_sexo.html')


@app.route('/api/perfil_paciente/sexo')
def api_perfil_paciente_sexo():
    """Retorna contagem de pacientes por sexo para a data selecionada (ou última disponível)."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        selected_date = request.args.get('data')
        with engine.connect() as conn:
            if not selected_date:
                sql_last = text("SELECT data_referencia FROM historico_ocupacao_completo ORDER BY data_referencia DESC LIMIT 1")
                result = conn.execute(sql_last).scalar()
                if result:
                    selected_date = result
                else:
                    return {"error": "Sem dados disponíveis"}, 404

            sql = text("""
                SELECT COALESCE(NULLIF(TRIM(sexo), ''), 'Não informado') as sexo,
                       COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE data_referencia = :data
                GROUP BY sexo
                ORDER BY cnt DESC
            """)
            rows = conn.execute(sql, {"data": selected_date}).mappings().all()

            labels = [r['sexo'] for r in rows]
            data = [int(r['cnt']) for r in rows]

            return {"labels": labels, "data": data}
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/perfil_paciente/idade')
def perfil_paciente_idade():
    return render_template('perfil_paciente_idade.html')


@app.route('/api/perfil_paciente/idade')
def api_perfil_paciente_idade():
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    try:
        selected_date = request.args.get('data')
        with engine.connect() as conn:
            if not selected_date:
                sql_last = text("SELECT data_referencia FROM historico_ocupacao_completo ORDER BY data_referencia DESC LIMIT 1")
                result = conn.execute(sql_last).scalar()
                if result:
                    selected_date = result
                else:
                    return {"error": "Sem dados disponíveis"}, 404

            sql = text("""
                SELECT
                  CASE
                    WHEN idade IS NULL THEN 'Não informado'
                    WHEN idade < 18 THEN '0-17'
                    WHEN idade BETWEEN 18 AND 29 THEN '18-29'
                    WHEN idade BETWEEN 30 AND 44 THEN '30-44'
                    WHEN idade BETWEEN 45 AND 59 THEN '45-59'
                    WHEN idade BETWEEN 60 AND 74 THEN '60-74'
                    ELSE '75+'
                  END as faixa,
                  COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE data_referencia = :data
                GROUP BY faixa
                ORDER BY FIELD(faixa, '0-17','18-29','30-44','45-59','60-74','75+','Não informado')
            """)
            rows = conn.execute(sql, {"data": selected_date}).mappings().all()
            labels = [r['faixa'] for r in rows]
            data = [int(r['cnt']) for r in rows]
            return {"labels": labels, "data": data}
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/perfil_paciente/tempo_permanencia')
def perfil_paciente_tempo():
    return render_template('perfil_paciente_tempo_permanencia.html')


@app.route('/api/perfil_paciente/tempo_permanencia')
def api_perfil_paciente_tempo():
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    try:
        selected_date = request.args.get('data')
        with engine.connect() as conn:
            if not selected_date:
                sql_last = text("SELECT data_referencia FROM historico_ocupacao_completo ORDER BY data_referencia DESC LIMIT 1")
                result = conn.execute(sql_last).scalar()
                if result:
                    selected_date = result
                else:
                    return {"error": "Sem dados disponíveis"}, 404

            sql = text("""
                SELECT
                  CASE
                    WHEN dias IS NULL THEN 'Não informado'
                    WHEN dias BETWEEN 0 AND 2 THEN '0-2'
                    WHEN dias BETWEEN 3 AND 7 THEN '3-7'
                    WHEN dias BETWEEN 8 AND 14 THEN '8-14'
                    WHEN dias BETWEEN 15 AND 30 THEN '15-30'
                    ELSE '31+'
                  END as faixa,
                  COUNT(*) as cnt
                FROM (
                  SELECT TIMESTAMPDIFF(DAY, data_internacao_leito, data_referencia) as dias
                  FROM historico_ocupacao_completo
                  WHERE data_referencia = :data
                ) as t
                GROUP BY faixa
                ORDER BY FIELD(faixa, '0-2','3-7','8-14','15-30','31+','Não informado')
            """)
            rows = conn.execute(sql, {"data": selected_date}).mappings().all()
            labels = [r['faixa'] for r in rows]
            data = [int(r['cnt']) for r in rows]
            return {"labels": labels, "data": data}
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/perfil_paciente/clinica')
def perfil_paciente_clinica():
    return render_template('perfil_paciente_clinica.html')


@app.route('/api/perfil_paciente/clinica')
def api_perfil_paciente_clinica():
    if not db_status:
        return {"error": "Banco não conectado"}, 500
    try:
        selected_date = request.args.get('data')
        with engine.connect() as conn:
            if not selected_date:
                sql_last = text("SELECT data_referencia FROM historico_ocupacao_completo ORDER BY data_referencia DESC LIMIT 1")
                result = conn.execute(sql_last).scalar()
                if result:
                    selected_date = result
                else:
                    return {"error": "Sem dados disponíveis"}, 404

            sql = text("""
                SELECT nome_enfermaria as clinica,
                       COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE data_referencia = :data
                GROUP BY nome_enfermaria
                ORDER BY cnt DESC
                LIMIT 50
            """)
            rows = conn.execute(sql, {"data": selected_date}).mappings().all()
            labels = [r['clinica'] for r in rows]
            data = [int(r['cnt']) for r in rows]
            return {"labels": labels, "data": data}
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/disponibilidade')
def api_disponibilidade():
    """Retorna série temporal de disponibilidade (vagos, ocupados, cedidos, impedidos, reservados)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        predio = request.args.get('predio')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')

        with engine.connect() as conn:
            where_conditions = []
            params = {}

            if periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            else:
                # If month filter is provided without explicit period, filter by the same year
                # as the latest available date and by the requested month. Otherwise default
                # to last 14 days window.
                sql_last = text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")
                last = conn.execute(sql_last).scalar()
                if not last:
                    return {"error": "Sem dados disponíveis"}, 404

                if mes:
                    year, month = _parse_mes_param(conn, mes)
                    if year is None or month is None:
                        return {"error": "Filtro de mes invalido"}, 400
                    params['mes'] = month
                    params['ano'] = year
                    where_conditions.append("MONTH(data_referencia) = :mes")
                    where_conditions.append("YEAR(data_referencia) = :ano")
                else:
                    where_conditions.append("data_referencia BETWEEN DATE_SUB(:last, INTERVAL 13 DAY) AND :last")
                    params['last'] = last
            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica
            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")

            where_clause = " AND ".join(where_conditions) if where_conditions else "1=1"

            sql = text(f"""
                SELECT DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COALESCE(SUM(CASE WHEN status_leito = 'LIVRE' THEN 1 ELSE 0 END), 0) as vagos,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'CEDIDO' THEN 1 ELSE 0 END), 0) as cedidos,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'RESERVADO' THEN 1 ELSE 0 END), 0) as reservados
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql, params).mappings().all()

            return {
                "labels": [r['dia'] for r in rows],
                "datasets": [
                    {"label": "Vagos", "data": [int(r['vagos']) for r in rows], "backgroundColor": "rgba(34,197,94,0.6)", "borderColor": "rgba(34,197,94,1)", "type": "bar", "stack": "status"},
                    {"label": "Ocupados", "data": [int(r['ocupados']) for r in rows], "backgroundColor": "rgba(59,130,246,0.6)", "borderColor": "rgba(59,130,246,1)", "type": "bar", "stack": "status"},
                    {"label": "Cedido", "data": [int(r['cedidos']) for r in rows], "backgroundColor": "rgba(255,159,64,0.6)", "borderColor": "rgba(255,159,64,1)", "type": "bar", "stack": "status"},
                    {"label": "Impedidos", "data": [int(r['impedidos']) for r in rows], "backgroundColor": "rgba(220,38,38,0.6)", "borderColor": "rgba(220,38,38,1)", "type": "bar", "stack": "status"},
                    {"label": "Reservados", "data": [int(r['reservados']) for r in rows], "backgroundColor": "rgba(148,163,184,0.6)", "borderColor": "rgba(148,163,184,1)", "type": "bar", "stack": "status"}
                ]
            }
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/tempo_permanencia')
def api_tempo_permanencia():
    """Retorna métricas e lista de pacientes por tempo de permanência"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        # filtros
        selected_date = request.args.get('data_referencia')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')
        predio = request.args.get('predio')
        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 50))

        with engine.connect() as conn:
            # Determina data de referência baseado nos filtros (prioridade: mes > periodo > data única)
            reference_date = None
            
            if mes:
                # Filtro por mês: pega última data do mês
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                sql_month_last = text("""
                    SELECT MAX(data_referencia) 
                    FROM historico_ocupacao_completo 
                    WHERE MONTH(data_referencia) = :mes AND YEAR(data_referencia) = :ano
                """)
                reference_date = conn.execute(sql_month_last, {"mes": month, "ano": year}).scalar()
            elif periodo_inicio and periodo_fim:
                # Filtro por período: usa data_fim como referência
                reference_date = periodo_fim
            elif selected_date:
                # Data única informada
                reference_date = selected_date
            else:
                # Padrão: última data disponível
                sql_last = text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")
                reference_date = conn.execute(sql_last).scalar()
            
            if not reference_date:
                return {"error": "Sem dados disponíveis"}, 404

            # Monta filtros WHERE para a seleção de pacientes ocupados
            where_conditions = ["status_leito = 'OCUPADO'"]
            params = {"data_referencia": reference_date}

            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica

            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")

            # período filtrado por data_referencia (aplica quando informado)
            if mes:
                year, month = _parse_mes_param(conn, mes)
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            elif periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            else:
                where_conditions.append("data_referencia = :data_referencia")

            where_clause = " AND ".join(where_conditions)

            # Consulta por paciente (agrupa por identificador: prioriza prontuario se não vazio/nulo)
            sql_patients = text(f"""
                SELECT
                    COALESCE(NULLIF(TRIM(prontuario), ''), cns_paciente, nome_paciente) as patient_id,
                    MIN(data_internacao) as data_internacao,
                    MAX(nome_paciente) as nome_paciente,
                    MAX(prontuario) as prontuario,
                    MAX(idade) as idade,
                    MAX(sexo) as sexo,
                    MAX(nome_enfermaria) as nome_enfermaria,
                    TIMESTAMPDIFF(DAY, MIN(data_internacao), :data_referencia) as dias
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY patient_id
                ORDER BY dias DESC
            """)

            rows = conn.execute(sql_patients, params).mappings().all()

            # Transforma em lista python
            patients = []
            dias_list = []
            for r in rows:
                dias = int(r['dias']) if r['dias'] is not None else 0
                dias_list.append(dias)
                patients.append({
                    'patient_id': r['patient_id'],
                    'nome': r['nome_paciente'] or '',
                    'prontuario': r['prontuario'],
                    'idade': int(r['idade']) if r['idade'] is not None else None,
                    'sexo': r['sexo'],
                    'clinica': r['nome_enfermaria'],
                    'data_internacao': str(r['data_internacao']) if r['data_internacao'] is not None else None,
                    'dias': dias
                })

            total = len(patients)
            avg_los = round(sum(dias_list) / total, 1) if total > 0 else 0

            # Mediana em Python
            median_los = 0
            if total > 0:
                sorted_days = sorted(dias_list)
                mid = total // 2
                if total % 2 == 0:
                    median_los = int((sorted_days[mid - 1] + sorted_days[mid]) / 2)
                else:
                    median_los = int(sorted_days[mid])

            # Counters
            long_gt_30 = sum(1 for d in dias_list if d > 30)
            long_gt_30_60 = sum(1 for p, d in zip(patients, dias_list) if d > 30 and (p.get('idade') or 0) >= 60)
            long_gt_30_ped = sum(1 for p, d in zip(patients, dias_list) if d > 30 and (p.get('idade') is not None and p.get('idade') < 18))

            # Histogram buckets
            buckets = {'0-7': 0, '8-14': 0, '15-30': 0, '31-60': 0, '61-90': 0, '>90': 0}
            for d in dias_list:
                if d <= 7:
                    buckets['0-7'] += 1
                elif d <= 14:
                    buckets['8-14'] += 1
                elif d <= 30:
                    buckets['15-30'] += 1
                elif d <= 60:
                    buckets['31-60'] += 1
                elif d <= 90:
                    buckets['61-90'] += 1
                else:
                    buckets['>90'] += 1

            # Paginação (base para KPIs/histograma)
            start = (page - 1) * per_page
            end = start + per_page
            page_items = patients[start:end]

            # Se houver filtros (parâmetros de filtros presentes), a tabela de 'Longa Permanência'
            # deve listar apenas pacientes com mais de 30 dias. Construímos uma segunda query
            # com HAVING para obter somente esses pacientes e paginamos sobre ela.
            filter_present = any([
                request.args.get('data_referencia'), request.args.get('periodo_inicio'), request.args.get('periodo_fim'),
                request.args.get('mes'), request.args.get('clinica'), request.args.get('predio')
            ])

            patients_table = patients
            patients_table_total = len(patients)

            if filter_present:
                # Garante que params contenha data_referencia para o TIMESTAMPDIFF
                params_longa = dict(params)
                if 'data_referencia' not in params_longa:
                    params_longa['data_referencia'] = selected_date

                sql_patients_longa = text(f"""
                    SELECT
                        COALESCE(NULLIF(TRIM(prontuario), ''), cns_paciente, nome_paciente) as patient_id,
                        MIN(data_internacao) as data_internacao,
                        MAX(nome_paciente) as nome_paciente,
                        MAX(prontuario) as prontuario,
                        MAX(idade) as idade,
                        MAX(sexo) as sexo,
                        MAX(nome_enfermaria) as nome_enfermaria,
                        TIMESTAMPDIFF(DAY, MIN(data_internacao), :data_referencia) as dias
                    FROM historico_ocupacao_completo
                    WHERE {where_clause}
                    GROUP BY patient_id
                    HAVING TIMESTAMPDIFF(DAY, MIN(data_internacao), :data_referencia) > 30
                    ORDER BY dias DESC
                """)

                rows_longa = conn.execute(sql_patients_longa, params_longa).mappings().all()
                patients_longa = []
                dias_list_longa = []
                for r in rows_longa:
                    dias = int(r['dias']) if r['dias'] is not None else 0
                    dias_list_longa.append(dias)
                    patients_longa.append({
                        'patient_id': r['patient_id'],
                        'nome': r['nome_paciente'] or '',
                        'prontuario': r['prontuario'],
                        'idade': int(r['idade']) if r['idade'] is not None else None,
                        'sexo': r['sexo'],
                        'clinica': r['nome_enfermaria'],
                        'data_internacao': str(r['data_internacao']) if r['data_internacao'] is not None else None,
                        'dias': dias
                    })

                patients_table = patients_longa
                patients_table_total = len(patients_longa)

                # Paginação sobre lista de longa permanência
                page_items = patients_table[(page-1)*per_page : (page-1)*per_page + per_page]

            # Mask names for frontend display (keep full name available only for export endpoint)
            def mask_name(full):
                if not full:
                    return ''
                parts = full.split()
                if len(parts) == 1:
                    s = parts[0]
                    if len(s) <= 2:
                        return s[0] + '*'
                    return s[0] + '*'*(len(s)-2) + s[-1]
                # first name + last initial
                first = parts[0]
                last = parts[-1]
                return f"{first} {last[0]}."

            for it in page_items:
                it['nome_masked'] = mask_name(it['nome'])

            response_data = {
                'data_referencia': str(reference_date),
                'total_patients': total,
                'avg_los': avg_los,
                'median_los': median_los,
                'long_gt_30': long_gt_30,
                'long_gt_30_60': long_gt_30_60,
                'long_gt_30_ped': long_gt_30_ped,
                'histogram': buckets,
                'page': page,
                'per_page': per_page,
                'patients': page_items,
                'patients_table_total': patients_table_total
            }
            
            # Adiciona informações sobre filtros aplicados
            if mes:
                response_data['filtro_aplicado'] = 'mes'
                response_data['mes'] = mes
            elif periodo_inicio and periodo_fim:
                response_data['filtro_aplicado'] = 'periodo'
                response_data['periodo_inicio'] = periodo_inicio
                response_data['periodo_fim'] = periodo_fim
            else:
                response_data['filtro_aplicado'] = 'data_unica'
            
            return jsonify(response_data)

    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/tempo_permanencia/export')
def api_tempo_permanencia_export():
    """Exporta a lista completa de longa permanência (>30 dias) em Excel (full names)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        selected_date = request.args.get('data_referencia')
        periodo_inicio = request.args.get('periodo_inicio')
        periodo_fim = request.args.get('periodo_fim')
        mes = request.args.get('mes')
        clinica = request.args.get('clinica')
        predio = request.args.get('predio')

        with engine.connect() as conn:
            # Determina data de referência (mesma lógica que api_tempo_permanencia)
            reference_date = None
            
            if mes:
                year, month = _parse_mes_param(conn, mes)
                if year is None or month is None:
                    return {"error": "Filtro de mes invalido"}, 400
                sql_month_last = text("""
                    SELECT MAX(data_referencia) 
                    FROM historico_ocupacao_completo 
                    WHERE MONTH(data_referencia) = :mes AND YEAR(data_referencia) = :ano
                """)
                reference_date = conn.execute(sql_month_last, {"mes": month, "ano": year}).scalar()
            elif periodo_inicio and periodo_fim:
                reference_date = periodo_fim
            elif selected_date:
                reference_date = selected_date
            else:
                sql_last = text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")
                reference_date = conn.execute(sql_last).scalar()
            
            if not reference_date:
                return {"error": "Sem dados disponíveis"}, 404

            where_conditions = ["status_leito = 'OCUPADO'"]
            params = {"data_referencia": reference_date}
            
            if clinica:
                where_conditions.append("nome_enfermaria = :clinica")
                params['clinica'] = clinica
            if predio == '1':
                where_conditions.append("num_enf BETWEEN 111 AND 199")
            elif predio == '2':
                where_conditions.append("num_enf BETWEEN 200 AND 299")
            
            # Aplica filtro de período
            if mes:
                year, month = _parse_mes_param(conn, mes)
                where_conditions.append("MONTH(data_referencia) = :mes")
                where_conditions.append("YEAR(data_referencia) = :ano")
                params['mes'] = month
                params['ano'] = year
            elif periodo_inicio and periodo_fim:
                where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
                params['periodo_inicio'] = periodo_inicio
                params['periodo_fim'] = periodo_fim
            else:
                where_conditions.append("data_referencia = :data_referencia")

            where_clause = " AND ".join(where_conditions)

            sql_export = text(f"""
                SELECT
                    COALESCE(NULLIF(TRIM(prontuario), ''), cns_paciente, nome_paciente) as patient_id,
                    MIN(data_internacao) as data_internacao,
                    MAX(nome_paciente) as nome_paciente,
                    MAX(prontuario) as prontuario,
                    MAX(idade) as idade,
                    MAX(sexo) as sexo,
                    MAX(nome_enfermaria) as nome_enfermaria,
                    TIMESTAMPDIFF(DAY, MIN(data_internacao), :data_referencia) as dias
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY patient_id
                HAVING dias > 30
                ORDER BY dias DESC
            """)

            rows = conn.execute(sql_export, params).mappings().all()
            # Converte para DataFrame
            df = pd.DataFrame([{
                'Nome': r['nome_paciente'],
                'Prontuario': r['prontuario'],
                'Idade': int(r['idade']) if r['idade'] is not None else None,
                'Sexo': r['sexo'],
                'Clinica': r['nome_enfermaria'],
                'Data Internacao': r['data_internacao'],
                'Dias Internado': int(r['dias'])
            } for r in rows])

            output = BytesIO()
            # Escreve Excel
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='LongaPermanencia')
            output.seek(0)

            filename = f"longa_permanencia_{selected_date}.xlsx"
            return send_file(output, download_name=filename, as_attachment=True, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    except Exception as e:
        return {"error": str(e)}, 500


def _parse_mes_param(conn, mes):
    if not mes:
        return None, None
    try:
        if isinstance(mes, str) and '-' in mes:
            parts = mes.split('-')
            if len(parts) >= 2:
                return int(parts[0]), int(parts[1])
        month = int(mes)
        year = conn.execute(text("SELECT YEAR(MAX(data_referencia)) FROM historico_ocupacao_completo")).scalar()
        return int(year) if year else None, month
    except Exception:
        return None, None


def _get_perfil_snapshot_context(conn, args):
    predio = args.get('predio')
    clinica = args.get('clinica')
    periodo_inicio = args.get('periodo_inicio')
    periodo_fim = args.get('periodo_fim')
    mes = args.get('mes')

    base_conditions = []
    base_params = {}

    if clinica:
        base_conditions.append("nome_enfermaria = :clinica")
        base_params['clinica'] = clinica

    if predio == '1':
        base_conditions.append("num_enf BETWEEN 111 AND 199")
    elif predio == '2':
        base_conditions.append("num_enf BETWEEN 200 AND 299")

    selected_date = None

    if periodo_fim:
        selected_date = periodo_fim
    elif periodo_inicio:
        selected_date = periodo_inicio
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is not None and month is not None:
            month_conditions = list(base_conditions)
            month_conditions.append("MONTH(data_referencia) = :mes")
            month_conditions.append("YEAR(data_referencia) = :ano")
            month_params = dict(base_params)
            month_params['mes'] = month
            month_params['ano'] = year
            month_where = " AND ".join(month_conditions) if month_conditions else "1=1"
            sql_month_last = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {month_where}")
            selected_date = conn.execute(sql_month_last, month_params).scalar()

    if not selected_date:
        latest_conditions = list(base_conditions)
        latest_where = " AND ".join(latest_conditions) if latest_conditions else "1=1"
        sql_last = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {latest_where}")
        selected_date = conn.execute(sql_last, base_params).scalar()

    if not selected_date:
        return None

    snapshot_conditions = list(base_conditions)
    snapshot_conditions.append("data_referencia = :data_referencia")
    snapshot_conditions.append("status_leito = 'OCUPADO'")
    snapshot_params = dict(base_params)
    snapshot_params['data_referencia'] = selected_date

    return {
        "selected_date": selected_date,
        "where": " AND ".join(snapshot_conditions),
        "params": snapshot_params
    }


def _get_perfil_range_context(conn, args):
    predio = args.get('predio')
    clinica = args.get('clinica')
    periodo_inicio = args.get('periodo_inicio')
    periodo_fim = args.get('periodo_fim')
    mes = args.get('mes')

    where_conditions = ["status_leito = 'OCUPADO'"]
    params = {}

    if clinica:
        where_conditions.append("nome_enfermaria = :clinica")
        params['clinica'] = clinica

    if predio == '1':
        where_conditions.append("num_enf BETWEEN 111 AND 199")
    elif predio == '2':
        where_conditions.append("num_enf BETWEEN 200 AND 299")

    if periodo_inicio and periodo_fim:
        where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
        params['periodo_inicio'] = periodo_inicio
        params['periodo_fim'] = periodo_fim
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is None or month is None:
            return None
        where_conditions.append("MONTH(data_referencia) = :mes")
        where_conditions.append("YEAR(data_referencia) = :ano")
        params['mes'] = month
        params['ano'] = year
    else:
        max_date = conn.execute(text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")).scalar()
        if not max_date:
            return None
        where_conditions.append("data_referencia BETWEEN DATE_SUB(:max_date, INTERVAL 13 DAY) AND :max_date")
        params['max_date'] = max_date

    return {
        "where": " AND ".join(where_conditions),
        "params": params
    }


def _get_perfil_range_context_all(conn, args):
    predio = args.get('predio')
    clinica = args.get('clinica')
    periodo_inicio = args.get('periodo_inicio')
    periodo_fim = args.get('periodo_fim')
    mes = args.get('mes')

    where_conditions = []
    params = {}

    if clinica:
        where_conditions.append("nome_enfermaria = :clinica")
        params['clinica'] = clinica

    if predio == '1':
        where_conditions.append("num_enf BETWEEN 111 AND 199")
    elif predio == '2':
        where_conditions.append("num_enf BETWEEN 200 AND 299")

    if periodo_inicio and periodo_fim:
        where_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
        params['periodo_inicio'] = periodo_inicio
        params['periodo_fim'] = periodo_fim
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is None or month is None:
            return None
        where_conditions.append("MONTH(data_referencia) = :mes")
        where_conditions.append("YEAR(data_referencia) = :ano")
        params['mes'] = month
        params['ano'] = year
    else:
        max_date = conn.execute(text("SELECT MAX(data_referencia) FROM historico_ocupacao_completo")).scalar()
        if not max_date:
            return None
        where_conditions.append("data_referencia BETWEEN DATE_SUB(:max_date, INTERVAL 13 DAY) AND :max_date")
        params['max_date'] = max_date

    return {
        "where": " AND ".join(where_conditions) if where_conditions else "1=1",
        "params": params
    }


def _patient_profile_query(where_clause):
    return f"""
        SELECT
            patient_id,
            nome_paciente,
            prontuario,
            cns_paciente,
            idade,
            sexo,
            nome_enfermaria,
            dias
        FROM (
            SELECT
                COALESCE(
                    NULLIF(TRIM(cns_paciente), ''),
                    NULLIF(TRIM(prontuario), ''),
                    NULLIF(TRIM(aih_paciente), ''),
                    NULLIF(TRIM(nome_paciente), ''),
                    CONCAT('LEITO-', num_enf, '-', leito)
                ) as patient_id,
                MAX(nome_paciente) as nome_paciente,
                MAX(prontuario) as prontuario,
                MAX(cns_paciente) as cns_paciente,
                MAX(CAST(idade AS SIGNED)) as idade,
                MAX(sexo) as sexo,
                MAX(nome_enfermaria) as nome_enfermaria,
                TIMESTAMPDIFF(DAY, MIN(data_internacao), :ref_date) as dias
            FROM historico_ocupacao_completo
            WHERE {where_clause}
            GROUP BY patient_id
        ) perfil
    """


@app.route('/api/perfil_paciente/resumo')
def api_perfil_paciente_resumo():
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            context = _get_perfil_snapshot_context(conn, request.args)
            if not context:
                return {"error": "Sem dados disponíveis"}, 404

            periodo_inicio = request.args.get('periodo_inicio')
            periodo_fim = request.args.get('periodo_fim')
            mes = request.args.get('mes')

            if (periodo_inicio and periodo_fim) or mes:
                range_context = _get_perfil_range_context(conn, request.args)
                if not range_context:
                    return {"error": "Sem dados disponíveis"}, 404
                params = dict(range_context['params'])
                params['ref_date'] = context['selected_date']
                sql = text(_patient_profile_query(range_context['where']))
            else:
                params = dict(context['params'])
                params['ref_date'] = context['selected_date']
                sql = text(_patient_profile_query(context['where']))

            rows = conn.execute(sql, params).mappings().all()

            total = len(rows)
            if total == 0:
                return jsonify({
                    "data_referencia": str(context['selected_date']),
                    "total_pacientes": 0,
                    "sexo_m_pct": 0,
                    "sexo_f_pct": 0,
                    "idade_media": 0,
                    "idade_mediana": 0,
                    "tempo_medio": 0,
                    "tempo_mediano": 0,
                    "longa_permanencia_pct": 0
                })

            idades = [int(r['idade']) for r in rows if r['idade'] is not None]
            tempos = [int(r['dias']) if r['dias'] is not None else 0 for r in rows]

            sex_norm = []
            for r in rows:
                sexo_raw = (r['sexo'] or '').strip().upper()
                if sexo_raw in ('M', 'MASCULINO'):
                    sex_norm.append('M')
                elif sexo_raw in ('F', 'FEMININO'):
                    sex_norm.append('F')
                else:
                    sex_norm.append('NI')

            def median(values):
                if not values:
                    return 0
                ordered = sorted(values)
                n = len(ordered)
                m = n // 2
                if n % 2 == 0:
                    return round((ordered[m - 1] + ordered[m]) / 2, 1)
                return ordered[m]

            male_count = sum(1 for s in sex_norm if s == 'M')
            female_count = sum(1 for s in sex_norm if s == 'F')
            longa_count = sum(1 for d in tempos if d > 30)

            return jsonify({
                "data_referencia": str(context['selected_date']),
                "total_pacientes": total,
                "sexo_m_pct": round((male_count / total) * 100, 1),
                "sexo_f_pct": round((female_count / total) * 100, 1),
                "idade_media": round(sum(idades) / len(idades), 1) if idades else 0,
                "idade_mediana": median(idades),
                "tempo_medio": round(sum(tempos) / len(tempos), 1) if tempos else 0,
                "tempo_mediano": median(tempos),
                "longa_permanencia_pct": round((longa_count / total) * 100, 1)
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/graficos')
def api_perfil_paciente_graficos():
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            snapshot_context = _get_perfil_snapshot_context(conn, request.args)
            if not snapshot_context:
                return {"error": "Sem dados disponíveis"}, 404

            periodo_inicio = request.args.get('periodo_inicio')
            periodo_fim = request.args.get('periodo_fim')
            mes = request.args.get('mes')

            if (periodo_inicio and periodo_fim) or mes:
                range_context = _get_perfil_range_context(conn, request.args)
                if not range_context:
                    return {"error": "Sem dados disponíveis"}, 404
                profile_params = dict(range_context['params'])
                profile_params['ref_date'] = snapshot_context['selected_date']
                sql_profile = text(_patient_profile_query(range_context['where']))
            else:
                profile_params = dict(snapshot_context['params'])
                profile_params['ref_date'] = snapshot_context['selected_date']
                sql_profile = text(_patient_profile_query(snapshot_context['where']))

            rows = conn.execute(sql_profile, profile_params).mappings().all()

            range_context = _get_perfil_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_series = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(DISTINCT COALESCE(
                        NULLIF(TRIM(cns_paciente), ''),
                        NULLIF(TRIM(prontuario), ''),
                        NULLIF(TRIM(aih_paciente), ''),
                        NULLIF(TRIM(nome_paciente), ''),
                        CONCAT('LEITO-', num_enf, '-', leito)
                    )) as pacientes_ativos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            series_rows = conn.execute(sql_series, range_context['params']).mappings().all()

            sexo_counts = {'Masculino': 0, 'Feminino': 0, 'Não informado': 0}
            faixa_counts = {'0-17': 0, '18-39': 0, '40-59': 0, '60-79': 0, '80+': 0}
            hist_counts = {'0-7': 0, '8-14': 0, '15-30': 0, '31-60': 0, '61-90': 0, '>90': 0}
            clinica_counts = {}
            clinica_los = {}

            for row in rows:
                sexo_raw = (row['sexo'] or '').strip().upper()
                if sexo_raw in ('M', 'MASCULINO'):
                    sexo_counts['Masculino'] += 1
                elif sexo_raw in ('F', 'FEMININO'):
                    sexo_counts['Feminino'] += 1
                else:
                    sexo_counts['Não informado'] += 1

                idade = int(row['idade']) if row['idade'] is not None else None
                if idade is not None:
                    if idade <= 17:
                        faixa_counts['0-17'] += 1
                    elif idade <= 39:
                        faixa_counts['18-39'] += 1
                    elif idade <= 59:
                        faixa_counts['40-59'] += 1
                    elif idade <= 79:
                        faixa_counts['60-79'] += 1
                    else:
                        faixa_counts['80+'] += 1

                dias = int(row['dias']) if row['dias'] is not None else 0
                if dias <= 7:
                    hist_counts['0-7'] += 1
                elif dias <= 14:
                    hist_counts['8-14'] += 1
                elif dias <= 30:
                    hist_counts['15-30'] += 1
                elif dias <= 60:
                    hist_counts['31-60'] += 1
                elif dias <= 90:
                    hist_counts['61-90'] += 1
                else:
                    hist_counts['>90'] += 1

                clinica = row['nome_enfermaria'] or 'Sem clínica'
                clinica_counts[clinica] = clinica_counts.get(clinica, 0) + 1
                if clinica not in clinica_los:
                    clinica_los[clinica] = {'sum': 0, 'count': 0}
                clinica_los[clinica]['sum'] += dias
                clinica_los[clinica]['count'] += 1

            top_clinicas = sorted(clinica_counts.items(), key=lambda item: item[1], reverse=True)[:10]
            clinica_avg_los = []
            for clinica, stats in clinica_los.items():
                avg_days = round(stats['sum'] / stats['count'], 1) if stats['count'] > 0 else 0
                clinica_avg_los.append((clinica, avg_days))
            top_clinicas_los = sorted(clinica_avg_los, key=lambda item: item[1], reverse=True)[:10]

            return jsonify({
                "serie_pacientes": {
                    "labels": [r['dia'] for r in series_rows],
                    "data": [int(r['pacientes_ativos']) for r in series_rows]
                },
                "sexo": {
                    "labels": list(sexo_counts.keys()),
                    "data": list(sexo_counts.values())
                },
                "faixa_etaria": {
                    "labels": list(faixa_counts.keys()),
                    "data": list(faixa_counts.values())
                },
                "hist_permanencia": {
                    "labels": list(hist_counts.keys()),
                    "data": list(hist_counts.values())
                },
                "top_clinicas": {
                    "labels": [c[0] for c in top_clinicas],
                    "data": [c[1] for c in top_clinicas]
                },
                "permanencia_clinica": {
                    "labels": [c[0] for c in top_clinicas_los],
                    "data": [c[1] for c in top_clinicas_los]
                }
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/tabela')
def api_perfil_paciente_tabela():
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        page = int(request.args.get('page', 1))
        per_page = int(request.args.get('per_page', 25))
        if page < 1:
            page = 1
        if per_page < 1:
            per_page = 25
        if per_page > 200:
            per_page = 200

        with engine.connect() as conn:
            context = _get_perfil_snapshot_context(conn, request.args)
            if not context:
                return {"error": "Sem dados disponíveis"}, 404

            periodo_inicio = request.args.get('periodo_inicio')
            periodo_fim = request.args.get('periodo_fim')
            mes = request.args.get('mes')

            if (periodo_inicio and periodo_fim) or mes:
                range_context = _get_perfil_range_context(conn, request.args)
                if not range_context:
                    return {"error": "Sem dados disponíveis"}, 404
                params = dict(range_context['params'])
                params['ref_date'] = context['selected_date']
                base_sql = _patient_profile_query(range_context['where'])
            else:
                params = dict(context['params'])
                params['ref_date'] = context['selected_date']
                base_sql = _patient_profile_query(context['where'])
            sql_count = text(f"SELECT COUNT(*) as total FROM ({base_sql}) x")
            total = int(conn.execute(sql_count, params).scalar() or 0)

            offset = (page - 1) * per_page
            sql_page = text(f"""
                {base_sql}
                ORDER BY dias DESC
                LIMIT :limit OFFSET :offset
            """)
            page_params = dict(params)
            page_params['limit'] = per_page
            page_params['offset'] = offset
            rows = conn.execute(sql_page, page_params).mappings().all()

            data = []
            for r in rows:
                data.append({
                    "patient_id": r['patient_id'],
                    "nome": r['nome_paciente'] or '',
                    "prontuario": r['prontuario'] or '',
                    "cns": r['cns_paciente'] or '',
                    "idade": int(r['idade']) if r['idade'] is not None else None,
                    "sexo": r['sexo'] or '',
                    "clinica": r['nome_enfermaria'] or '',
                    "dias": int(r['dias']) if r['dias'] is not None else 0
                })

            return jsonify({
                "data_referencia": str(context['selected_date']),
                "page": page,
                "per_page": per_page,
                "total": total,
                "rows": data
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/export')
def api_perfil_paciente_export():
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            context = _get_perfil_snapshot_context(conn, request.args)
            if not context:
                return {"error": "Sem dados disponíveis"}, 404

            params = dict(context['params'])
            params['ref_date'] = context['selected_date']
            sql = text(f"""
                {_patient_profile_query(context['where'])}
                ORDER BY dias DESC
            """)
            rows = conn.execute(sql, params).mappings().all()

            df = pd.DataFrame([{
                'Paciente': r['nome_paciente'],
                'Prontuario': r['prontuario'],
                'CNS': r['cns_paciente'],
                'Idade': int(r['idade']) if r['idade'] is not None else None,
                'Sexo': r['sexo'],
                'Clinica': r['nome_enfermaria'],
                'Dias Internado': int(r['dias']) if r['dias'] is not None else 0
            } for r in rows])

            output = BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='PerfilPaciente')
            output.seek(0)

            filename = f"perfil_paciente_{context['selected_date']}.xlsx"
            return send_file(
                output,
                download_name=filename,
                as_attachment=True,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/impedimentos-serie')
def api_perfil_paciente_impedimentos_serie():
    """Retorna série temporal da taxa de impedimentos por dia (%)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_impedimentos = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' THEN 1 ELSE 0 END), 0) as impedidos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql_impedimentos, range_context['params']).mappings().all()

            return jsonify({
                "labels": [r['dia'] for r in rows],
                "data": [round((int(r['impedidos']) / int(r['total'])) * 100, 1) if int(r['total']) > 0 else 0 for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/status-serie')
def api_perfil_paciente_status_serie():
    """Retorna série temporal de status de leitos (ocupado/livre/impedido/etc)"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_status = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COALESCE(SUM(CASE WHEN status_leito = 'LIVRE' THEN 1 ELSE 0 END), 0) as livres,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'CEDIDO' THEN 1 ELSE 0 END), 0) as cedidos,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'RESERVADO' THEN 1 ELSE 0 END), 0) as reservados
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql_status, range_context['params']).mappings().all()

            return jsonify({
                "labels": [r['dia'] for r in rows],
                "datasets": [
                    {"label": "Livres", "data": [int(r['livres']) for r in rows]},
                    {"label": "Ocupados", "data": [int(r['ocupados']) for r in rows]},
                    {"label": "Cedido", "data": [int(r['cedidos']) for r in rows]},
                    {"label": "Impedidos", "data": [int(r['impedidos']) for r in rows]},
                    {"label": "Reservados", "data": [int(r['reservados']) for r in rows]}
                ]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/impedimentos-top')
def api_perfil_paciente_impedimentos_top():
    """Retorna top motivos de impedimento no recorte"""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            where_clause = f"({range_context['where']}) AND (status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%')"
            sql_top = text(f"""
                SELECT COALESCE(NULLIF(TRIM(motivo_impedimento), ''), 'Sem motivo informado') as motivo,
                       COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE {where_clause}
                GROUP BY motivo
                ORDER BY cnt DESC
                LIMIT 10
            """)
            rows = conn.execute(sql_top, range_context['params']).mappings().all()

            return jsonify({
                "labels": [r['motivo'] for r in rows],
                "data": [int(r['cnt']) for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/kpis-complementares')
def api_perfil_paciente_kpis_complementares():
    """Retorna KPIs complementares do perfil do paciente."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_stats = text(f"""
                SELECT
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN status_leito LIKE '%IMPEDIDO%' OR status_leito LIKE '%BLOQUEADO%' THEN 1 ELSE 0 END), 0) as impedidos,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END), 0) as ocupados,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' AND UPPER(TRIM(cronico)) IN ('SIM','S','1','TRUE','T','YES') THEN 1 ELSE 0 END), 0) as cronicos,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' AND NULLIF(TRIM(modo_ventilatorio), '') IS NOT NULL THEN 1 ELSE 0 END), 0) as ventilacao,
                    COALESCE(SUM(CASE WHEN status_leito = 'OCUPADO' AND UPPER(TRIM(inserido_no_trs)) IN ('SIM','S','1','TRUE','T','YES') THEN 1 ELSE 0 END), 0) as trs
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
            """)
            stats = conn.execute(sql_stats, range_context['params']).mappings().fetchone()

            sql_top_enf = text(f"""
                SELECT COALESCE(NULLIF(TRIM(nome_enfermaria), ''), 'Sem clínica') as enfermaria,
                       COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
                GROUP BY enfermaria
                ORDER BY cnt DESC
                LIMIT 5
            """)
            top_rows = conn.execute(sql_top_enf, range_context['params']).mappings().all()

            total = int(stats['total'] or 0)
            ocupados = int(stats['ocupados'] or 0)
            impedidos = int(stats['impedidos'] or 0)
            cronicos = int(stats['cronicos'] or 0)
            ventilacao = int(stats['ventilacao'] or 0)
            trs = int(stats['trs'] or 0)

            return jsonify({
                "impedidos_pct": round((impedidos / total) * 100, 1) if total > 0 else 0,
                "top_enfermarias": [{"label": r['enfermaria'], "count": int(r['cnt'])} for r in top_rows],
                "cronicos_pct": round((cronicos / ocupados) * 100, 1) if ocupados > 0 else 0,
                "ventilacao_pct": round((ventilacao / ocupados) * 100, 1) if ocupados > 0 else 0,
                "trs_pct": round((trs / ocupados) * 100, 1) if ocupados > 0 else 0
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/kpis-avancados')
def api_perfil_paciente_kpis_avancados():
    """Retorna KPIs avancados do perfil do paciente."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_rot = text(f"""
                SELECT
                    COUNT(DISTINCT COALESCE(
                        NULLIF(TRIM(cns_paciente), ''),
                        NULLIF(TRIM(prontuario), ''),
                        NULLIF(TRIM(aih_paciente), ''),
                        NULLIF(TRIM(nome_paciente), ''),
                        CONCAT('LEITO-', num_enf, '-', leito)
                    )) as pacientes,
                    COUNT(DISTINCT CONCAT(num_enf, '-', leito)) as leitos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
            """)
            rot = conn.execute(sql_rot, range_context['params']).mappings().fetchone()

            sql_motivo = text(f"""
                SELECT
                    COUNT(*) as total,
                    COALESCE(SUM(CASE WHEN NULLIF(TRIM(situacao_motivo_permanencia), '') IS NOT NULL THEN 1 ELSE 0 END), 0) as com_motivo
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
            """)
            motivo = conn.execute(sql_motivo, range_context['params']).mappings().fetchone()

            sql_reserva = text(f"""
                SELECT
                    AVG(TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito)) as media_dias
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                  AND data_sol_reserva IS NOT NULL
                  AND data_internacao_leito IS NOT NULL
                  AND TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito) >= 0
            """)
            reserva = conn.execute(sql_reserva, range_context['params']).scalar()

            pacientes = int(rot['pacientes'] or 0)
            leitos = int(rot['leitos'] or 0)
            total = int(motivo['total'] or 0)
            com_motivo = int(motivo['com_motivo'] or 0)

            return jsonify({
                "rotatividade": round((pacientes / leitos), 2) if leitos > 0 else 0,
                "motivo_permanencia_pct": round((com_motivo / total) * 100, 1) if total > 0 else 0,
                "tempo_medio_reserva": round(float(reserva), 1) if reserva is not None else None
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/rotatividade-serie')
def api_perfil_paciente_rotatividade_serie():
    """Retorna serie temporal de rotatividade por leito (pacientes / leito)."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    COUNT(DISTINCT COALESCE(
                        NULLIF(TRIM(cns_paciente), ''),
                        NULLIF(TRIM(prontuario), ''),
                        NULLIF(TRIM(aih_paciente), ''),
                        NULLIF(TRIM(nome_paciente), ''),
                        CONCAT('LEITO-', num_enf, '-', leito)
                    )) as pacientes,
                    COUNT(DISTINCT CONCAT(num_enf, '-', leito)) as leitos
                FROM historico_ocupacao_completo
                WHERE {range_context['where']} AND status_leito = 'OCUPADO'
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql, range_context['params']).mappings().all()

            return jsonify({
                "labels": [r['dia'] for r in rows],
                "data": [round((int(r['pacientes']) / int(r['leitos'])), 2) if int(r['leitos']) > 0 else 0 for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/longa-permanencia-ranking')
def api_perfil_paciente_longa_permanencia_ranking():
    """Retorna ranking de longa permanencia por clinica."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            context = _get_perfil_snapshot_context(conn, request.args)
            if not context:
                return {"error": "Sem dados disponíveis"}, 404

            params = dict(context['params'])
            params['ref_date'] = context['selected_date']
            base_sql = _patient_profile_query(context['where'])
            sql = text(f"""
                SELECT nome_enfermaria as clinica,
                       COUNT(*) as cnt,
                       AVG(dias) as media
                FROM ({base_sql}) t
                WHERE dias > 30
                GROUP BY clinica
                ORDER BY cnt DESC
                LIMIT 10
            """)
            rows = conn.execute(sql, params).mappings().all()

            return jsonify({
                "labels": [r['clinica'] or 'Sem clinica' for r in rows],
                "data": [int(r['cnt']) for r in rows],
                "avg": [round(float(r['media'] or 0), 1) for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/reserva-serie')
def api_perfil_paciente_reserva_serie():
    """Retorna serie temporal do tempo medio entre reserva e ocupacao."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context_all(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql = text(f"""
                SELECT
                    DATE_FORMAT(data_referencia, '%d/%m') as dia,
                    AVG(TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito)) as media_dias
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                  AND data_sol_reserva IS NOT NULL
                  AND data_internacao_leito IS NOT NULL
                  AND TIMESTAMPDIFF(DAY, data_sol_reserva, data_internacao_leito) >= 0
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            rows = conn.execute(sql, range_context['params']).mappings().all()

            return jsonify({
                "labels": [r['dia'] for r in rows],
                "data": [round(float(r['media_dias'] or 0), 1) for r in rows]
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/api/perfil_paciente/ocupacao-heatmap')
def api_perfil_paciente_ocupacao_heatmap():
    """Retorna matriz de ocupacao por dia da semana x clinica (top 10 clinicas)."""
    if not db_status:
        return {"error": "Banco não conectado"}, 500

    try:
        with engine.connect() as conn:
            range_context = _get_perfil_range_context(conn, request.args)
            if not range_context:
                return {"error": "Sem dados disponíveis"}, 404

            sql_heatmap = text(f"""
                SELECT
                    COALESCE(NULLIF(TRIM(nome_enfermaria), ''), 'Sem clínica') as clinica,
                    DAYOFWEEK(data_referencia) as dow,
                    COUNT(*) as cnt
                FROM historico_ocupacao_completo
                WHERE {range_context['where']}
                GROUP BY clinica, dow
            """)
            rows = conn.execute(sql_heatmap, range_context['params']).mappings().all()

            totals = {}
            values = {}
            for r in rows:
                clinica = r['clinica']
                totals[clinica] = totals.get(clinica, 0) + int(r['cnt'])
                values.setdefault(clinica, {})[int(r['dow'])] = int(r['cnt'])

            top_clinicas = [c for c, _ in sorted(totals.items(), key=lambda x: x[1], reverse=True)[:10]]
            dias = ['Dom', 'Seg', 'Ter', 'Qua', 'Qui', 'Sex', 'Sab']

            matrix = []
            max_val = 0
            for clinica in top_clinicas:
                row_vals = []
                for i, _ in enumerate(dias, start=1):
                    val = values.get(clinica, {}).get(i, 0)
                    row_vals.append(val)
                    if val > max_val:
                        max_val = val
                matrix.append(row_vals)

            return jsonify({
                "clinicas": top_clinicas,
                "dias": dias,
                "values": matrix,
                "max": max_val
            })
    except Exception as e:
        return {"error": str(e)}, 500


@app.route('/perfil_paciente')
def perfil_paciente():
    return render_template('perfil_paciente.html')

# ROTA PARA O PAINEL (renderiza template estático)
@app.route('/painel')
def painel():
    return render_template('painel.html')

@app.route('/emergencia')
def emergencia():
    return render_template('emergencia.html')

@app.route('/tempo_permanencia')
def tempo_permanencia():
    return render_template('tempo_permanencia.html')


@app.route('/relatorios')
def relatorios():
    return render_template('relatorios.html')


def _build_relatorios_base_filters(conn, filters):
    predio = filters.get('predio')
    clinica = filters.get('clinica')
    periodo_inicio = filters.get('periodo_inicio')
    periodo_fim = filters.get('periodo_fim')
    mes = filters.get('mes')
    data_referencia = filters.get('data_referencia')

    base_conditions = []
    base_params = {}

    if clinica:
        base_conditions.append("nome_enfermaria = :clinica")
        base_params['clinica'] = clinica

    if predio == '1':
        base_conditions.append("num_enf BETWEEN 111 AND 199")
    elif predio == '2':
        base_conditions.append("num_enf BETWEEN 200 AND 299")

    selected_date = data_referencia

    if not selected_date:
        if periodo_fim:
            selected_date = periodo_fim
        elif periodo_inicio:
            selected_date = periodo_inicio
        elif mes:
            year, month = _parse_mes_param(conn, mes)
            if year is not None and month is not None:
                month_conditions = list(base_conditions)
                month_conditions.append("MONTH(data_referencia) = :mes")
                month_conditions.append("YEAR(data_referencia) = :ano")
                month_params = dict(base_params)
                month_params['mes'] = month
                month_params['ano'] = year
                month_where = " AND ".join(month_conditions) if month_conditions else "1=1"
                sql_month_last = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {month_where}")
                selected_date = conn.execute(sql_month_last, month_params).scalar()

    if not selected_date:
        latest_where = " AND ".join(base_conditions) if base_conditions else "1=1"
        sql_last = text(f"SELECT MAX(data_referencia) FROM historico_ocupacao_completo WHERE {latest_where}")
        selected_date = conn.execute(sql_last, base_params).scalar()

    snapshot_conditions = list(base_conditions)
    snapshot_params = dict(base_params)
    if selected_date:
        snapshot_conditions.append("data_referencia = :data_referencia")
        snapshot_params['data_referencia'] = selected_date

    range_conditions = list(base_conditions)
    range_params = dict(base_params)
    if periodo_inicio and periodo_fim:
        range_conditions.append("data_referencia BETWEEN :periodo_inicio AND :periodo_fim")
        range_params['periodo_inicio'] = periodo_inicio
        range_params['periodo_fim'] = periodo_fim
    elif mes:
        year, month = _parse_mes_param(conn, mes)
        if year is not None and month is not None:
            range_conditions.append("MONTH(data_referencia) = :mes")
            range_conditions.append("YEAR(data_referencia) = :ano")
            range_params['mes'] = month
            range_params['ano'] = year
    elif selected_date:
        range_conditions.append("data_referencia BETWEEN DATE_SUB(:data_referencia, INTERVAL 13 DAY) AND :data_referencia")
        range_params['data_referencia'] = selected_date

    return {
        "selected_date": selected_date,
        "snapshot_where": " AND ".join(snapshot_conditions) if snapshot_conditions else "1=1",
        "snapshot_params": snapshot_params,
        "range_where": " AND ".join(range_conditions) if range_conditions else "1=1",
        "range_params": range_params
    }


def _build_relatorios_payload(filters, selected_blocks):
    payload = {
        "generated_at": datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        "filters": filters,
        "kpis": [],
        "charts": [],
        "tables": []
    }

    allowed_blocks = {
        "kpi_ocupacao",
        "chart_ocupacao_clinica",
        "chart_evolucao_ocupacao",
        "table_longa_permanencia",
        "kpi_emergencia"
    }

    # Remove duplicados preservando ordem e aplica lista de permitidos
    seen = set()
    blocks = []
    for block in selected_blocks:
        if block in allowed_blocks and block not in seen:
            seen.add(block)
            blocks.append(block)

    if len(blocks) > MAX_RELATORIO_BLOCKS:
        blocks = blocks[:MAX_RELATORIO_BLOCKS]

    if not blocks:
        return payload

    with engine.connect() as conn:
        context = _build_relatorios_base_filters(conn, filters)

        if "kpi_ocupacao" in blocks:
            sql_kpi_ocupacao = text(f"""
                SELECT
                    SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END) AS ocupados,
                    SUM(CASE WHEN status_leito = 'VAGO' THEN 1 ELSE 0 END) AS vagos,
                    SUM(CASE WHEN status_leito = 'IMPEDIDO' THEN 1 ELSE 0 END) AS impedidos,
                    COUNT(*) AS total
                FROM historico_ocupacao_completo
                WHERE {context['snapshot_where']}
            """)
            kpi_row = conn.execute(sql_kpi_ocupacao, context['snapshot_params']).mappings().first()
            if kpi_row:
                payload["kpis"].extend([
                    {"id": "ocupados", "label": "Leitos Ocupados", "value": int(kpi_row['ocupados'] or 0)},
                    {"id": "vagos", "label": "Leitos Vagos", "value": int(kpi_row['vagos'] or 0)},
                    {"id": "impedidos", "label": "Leitos Impedidos", "value": int(kpi_row['impedidos'] or 0)},
                    {"id": "total", "label": "Total de Leitos", "value": int(kpi_row['total'] or 0)}
                ])

        if "chart_ocupacao_clinica" in blocks:
            sql_chart_clinica = text(f"""
                SELECT nome_enfermaria AS clinica, COUNT(*) AS qtd
                FROM historico_ocupacao_completo
                WHERE {context['snapshot_where']} AND status_leito = 'OCUPADO'
                GROUP BY nome_enfermaria
                ORDER BY qtd DESC
                LIMIT 10
            """)
            chart_rows = conn.execute(sql_chart_clinica, context['snapshot_params']).mappings().all()
            payload["charts"].append({
                "id": "chart_ocupacao_clinica",
                "title": "Ocupação por Clínica",
                "type": "bar",
                "labels": [r['clinica'] for r in chart_rows],
                "data": [int(r['qtd']) for r in chart_rows]
            })

        if "chart_evolucao_ocupacao" in blocks:
            sql_evolucao = text(f"""
                SELECT DATE_FORMAT(data_referencia, '%d/%m') AS dia, COUNT(*) AS qtd
                FROM historico_ocupacao_completo
                WHERE {context['range_where']} AND status_leito = 'OCUPADO'
                GROUP BY data_referencia
                ORDER BY data_referencia
            """)
            evo_rows = conn.execute(sql_evolucao, context['range_params']).mappings().all()
            payload["charts"].append({
                "id": "chart_evolucao_ocupacao",
                "title": "Evolução da Ocupação",
                "type": "line",
                "labels": [r['dia'] for r in evo_rows],
                "data": [int(r['qtd']) for r in evo_rows]
            })

        if "table_longa_permanencia" in blocks and context['selected_date']:
            sql_longa = text(f"""
                SELECT
                    IFNULL(NULLIF(prontuario, ''), nome_paciente) AS patient_id,
                    MAX(nome_paciente) AS nome,
                    MAX(nome_enfermaria) AS clinica,
                    MIN(data_internacao) AS data_internacao,
                    TIMESTAMPDIFF(DAY, MIN(data_internacao), :selected_date) AS dias
                FROM historico_ocupacao_completo
                WHERE {context['snapshot_where']} AND status_leito = 'OCUPADO' AND data_internacao IS NOT NULL
                GROUP BY patient_id
                HAVING dias > 30
                ORDER BY dias DESC
                LIMIT 100
            """)
            longa_params = dict(context['snapshot_params'])
            longa_params['selected_date'] = context['selected_date']
            longa_rows = conn.execute(sql_longa, longa_params).mappings().all()
            payload["tables"].append({
                "id": "table_longa_permanencia",
                "title": "Pacientes com Longa Permanência",
                "columns": ["Paciente", "Clínica", "Data Internação", "Dias Internado"],
                "rows": [[
                    r['nome'] or '—',
                    r['clinica'] or '—',
                    r['data_internacao'].strftime('%d/%m/%Y') if r['data_internacao'] else '—',
                    int(r['dias'] or 0)
                ] for r in longa_rows[:MAX_TABELA_ROWS]]
            })

        if "kpi_emergencia" in blocks:
            ward_list = "', '".join(EMERGENCY_WARDS)
            sql_kpi_emergencia = text(f"""
                SELECT
                    SUM(CASE WHEN status_leito = 'OCUPADO' THEN 1 ELSE 0 END) AS ocupados,
                    COUNT(*) AS total
                FROM historico_ocupacao_completo
                WHERE {context['snapshot_where']} AND nome_enfermaria IN ('{ward_list}')
            """)
            e_row = conn.execute(sql_kpi_emergencia, context['snapshot_params']).mappings().first()
            ocupados = int((e_row or {}).get('ocupados') or 0)
            total = int((e_row or {}).get('total') or 0)
            taxa = round((ocupados / total) * 100, 1) if total else 0
            payload["kpis"].extend([
                {"id": "emerg_ocupados", "label": "Emergência Ocupados", "value": ocupados},
                {"id": "emerg_total", "label": "Emergência Total", "value": total},
                {"id": "emerg_taxa", "label": "Taxa Emergência", "value": f"{taxa}%"}
            ])

    return payload


@app.route('/api/relatorios/preview', methods=['POST'])
def relatorios_preview():
    try:
        body = request.get_json() or {}
        selected_blocks = body.get('selected_blocks', [])
        filters = body.get('filters', {})

        if not isinstance(selected_blocks, list):
            return jsonify({"error": "selected_blocks deve ser uma lista"}), 400
        if not isinstance(filters, dict):
            return jsonify({"error": "filters deve ser um objeto"}), 400
        if len(selected_blocks) > MAX_RELATORIO_BLOCKS:
            return jsonify({"error": f"Máximo de {MAX_RELATORIO_BLOCKS} blocos por relatório"}), 400

        payload = _build_relatorios_payload(filters, selected_blocks)
        return jsonify(payload)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/relatorios/webhook-config', methods=['GET', 'POST'])
def relatorios_webhook_config():
    if request.method == 'GET':
        return jsonify(WEBHOOK_CONFIG)

    try:
        data = request.get_json() or {}
        url = (data.get('url') or '').strip()
        enabled = bool(data.get('enabled'))

        if url and not (url.startswith('http://') or url.startswith('https://')):
            return jsonify({"error": "URL do webhook deve iniciar com http:// ou https://"}), 400

        WEBHOOK_CONFIG['url'] = url
        WEBHOOK_CONFIG['enabled'] = enabled and bool(url)
        return jsonify({"message": "Configuração salva", "config": WEBHOOK_CONFIG})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/relatorios/webhook-test', methods=['POST'])
def relatorios_webhook_test():
    try:
        url = WEBHOOK_CONFIG.get('url', '').strip()
        if not url:
            return jsonify({"error": "Webhook não configurado"}), 400

        test_payload = {
            "event": "nir_webhook_test",
            "generated_at": datetime.now().isoformat(),
            "message": "Teste de conexão do NIR Dashboard"
        }
        response = requests.post(url, json=test_payload, timeout=WEBHOOK_TIMEOUT_SECONDS)
        return jsonify({
            "message": "Teste enviado",
            "status_code": response.status_code,
            "ok": response.ok
        })
    except requests.RequestException as e:
        return jsonify({"error": f"Falha ao conectar no webhook: {str(e)}"}), 502
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/relatorios/webhook-send', methods=['POST'])
def relatorios_webhook_send():
    try:
        if not WEBHOOK_CONFIG.get('enabled'):
            return jsonify({"error": "Webhook desabilitado. Ative nas configurações."}), 400

        url = WEBHOOK_CONFIG.get('url', '').strip()
        if not url:
            return jsonify({"error": "Webhook não configurado"}), 400

        body = request.get_json() or {}
        selected_blocks = body.get('selected_blocks', [])
        filters = body.get('filters', {})
        custom_message = body.get('message', '')

        report_payload = _build_relatorios_payload(filters, selected_blocks)
        envelope = {
            "event": "nir_relatorio_manual",
            "generated_at": datetime.now().isoformat(),
            "message": custom_message,
            "report": report_payload
        }

        response = requests.post(url, json=envelope, timeout=WEBHOOK_TIMEOUT_SECONDS)
        return jsonify({
            "message": "Dados enviados ao n8n",
            "status_code": response.status_code,
            "ok": response.ok
        })
    except requests.RequestException as e:
        return jsonify({"error": f"Falha ao enviar para webhook: {str(e)}"}), 502
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/export/pptx', methods=['POST'])
def export_pptx():
    try:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
        except Exception:
            return jsonify({"error": "Dependência python-pptx não instalada"}), 500

        data = request.get_json() or {}
        page_title = data.get('page_title', 'NIR Dashboard - Relatório')
        filters = data.get('filters', {})
        kpis = data.get('kpis', [])
        charts = data.get('charts', [])
        tables = data.get('tables', [])

        prs = Presentation()

        slide_title = prs.slides.add_slide(prs.slide_layouts[0])
        slide_title.shapes.title.text = page_title
        slide_title.placeholders[1].text = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"

        slide_filters = prs.slides.add_slide(prs.slide_layouts[5])
        slide_filters.shapes.title.text = "Filtros Aplicados"
        tf = slide_filters.shapes.add_textbox(Inches(0.7), Inches(1.4), Inches(12), Inches(4)).text_frame
        tf.word_wrap = True
        if filters:
            for key, val in filters.items():
                p = tf.add_paragraph()
                p.text = f"{key}: {val}"
                p.font.size = Pt(18)
        else:
            tf.text = "Sem filtros aplicados"

        if kpis:
            slide_kpi = prs.slides.add_slide(prs.slide_layouts[5])
            slide_kpi.shapes.title.text = "Indicadores"
            y = 1.4
            for kpi in kpis:
                box = slide_kpi.shapes.add_textbox(Inches(0.7), Inches(y), Inches(12), Inches(0.5))
                t = box.text_frame
                t.text = f"{kpi.get('label', '')}: {kpi.get('value', '—')}"
                t.paragraphs[0].font.size = Pt(20)
                y += 0.55
                if y > 6.8:
                    break

        for chart in charts:
            image_data = chart.get('image', '')
            if not image_data or ',' not in image_data:
                continue
            slide_chart = prs.slides.add_slide(prs.slide_layouts[5])
            slide_chart.shapes.title.text = chart.get('title', 'Gráfico')
            b64 = image_data.split(',', 1)[1]
            img_bytes = BytesIO(base64.b64decode(b64))
            slide_chart.shapes.add_picture(img_bytes, Inches(0.6), Inches(1.2), Inches(12.1), Inches(5.8))

        for table in tables:
            columns = table.get('columns', [])
            rows = table.get('rows', [])
            if not columns:
                continue

            slide_table = prs.slides.add_slide(prs.slide_layouts[5])
            slide_table.shapes.title.text = table.get('title', 'Tabela')

            max_rows = min(len(rows), 14)
            table_shape = slide_table.shapes.add_table(
                max_rows + 1,
                len(columns),
                Inches(0.4),
                Inches(1.2),
                Inches(12.5),
                Inches(5.8)
            )
            ppt_table = table_shape.table

            for col_idx, col_name in enumerate(columns):
                ppt_table.cell(0, col_idx).text = str(col_name)

            for row_idx, row_data in enumerate(rows[:max_rows], start=1):
                for col_idx in range(len(columns)):
                    value = row_data[col_idx] if col_idx < len(row_data) else ''
                    ppt_table.cell(row_idx, col_idx).text = str(value)

        output = BytesIO()
        prs.save(output)
        output.seek(0)

        filename = f"relatorio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        return send_file(
            output,
            download_name=filename,
            as_attachment=True,
            mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation'
        )
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/export/pdf', methods=['POST'])
def export_pdf():
    """Generate PDF report from dashboard data (KPIs + charts as base64 images)."""
    try:
        print("=== Iniciando geração de PDF ===", flush=True)
        data = request.get_json()
        if not data:
            print("Erro: Nenhum dado recebido", flush=True)
            return jsonify({"error": "No data provided"}), 400

        print(f"Dados recebidos: {len(str(data))} caracteres", flush=True)
        page_title = data.get('page_title', 'NIR Dashboard - Relatório')
        filters = data.get('filters', {})
        kpis = data.get('kpis', [])
        charts = data.get('charts', [])
        tables = data.get('tables', [])
        print(f"KPIs: {len(kpis)}, Gráficos: {len(charts)}, Tabelas: {len(tables)}", flush=True)

        page_title_safe = escape(str(page_title))
        
        # Build filter summary text
        filter_text = []
        if filters.get('predio'):
            filter_text.append(f"Prédio {filters['predio']}")
        if filters.get('clinica'):
            filter_text.append(f"Clínica: {filters['clinica']}")
        if filters.get('periodo_inicio') and filters.get('periodo_fim'):
            filter_text.append(f"Período: {filters['periodo_inicio']} a {filters['periodo_fim']}")
        elif filters.get('mes'):
            meses = ['', 'Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
                     'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
            filter_text.append(f"Mês: {meses[int(filters['mes'])]}")
        if filters.get('data_referencia'):
            filter_text.append(f"Data: {filters['data_referencia']}")
        
        filters_display = ' | '.join(filter_text) if filter_text else 'Todos os dados'
        
        # Build HTML report
        html_content = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="UTF-8">
    <title>{page_title_safe}</title>
    <style>
        @page {{
            size: A4;
            margin: 2cm;
        }}
        body {{
            font-family: Arial, sans-serif;
            color: #1f2937;
            line-height: 1.6;
        }}
        .header {{
            text-align: center;
            margin-bottom: 2rem;
            padding-bottom: 1rem;
            border-bottom: 3px solid #0b72d9;
        }}
        .header h1 {{
            color: #0b72d9;
            margin: 0;
            font-size: 24pt;
        }}
        .header .subtitle {{
            color: #6b7280;
            margin-top: 0.5rem;
            font-size: 10pt;
        }}
        .filters {{
            background: #f3f4f6;
            padding: 1rem;
            margin-bottom: 2rem;
            border-radius: 8px;
            font-size: 9pt;
        }}
        .filters strong {{
            color: #0b72d9;
        }}
        .kpi-grid {{
            display: grid;
            grid-template-columns: repeat(4, 1fr);
            gap: 1rem;
            margin-bottom: 2rem;
        }}
        .kpi-card {{
            background: #f9fafb;
            border-left: 4px solid #0b72d9;
            padding: 1rem;
            border-radius: 4px;
        }}
        .kpi-label {{
            font-size: 8pt;
            color: #6b7280;
            text-transform: uppercase;
            font-weight: 600;
            margin-bottom: 0.5rem;
        }}
        .kpi-value {{
            font-size: 20pt;
            font-weight: bold;
            color: #1f2937;
        }}
        .chart-section {{
            margin-bottom: 2rem;
            page-break-inside: avoid;
        }}
        .chart-title {{
            font-size: 12pt;
            font-weight: bold;
            color: #1f2937;
            margin-bottom: 1rem;
        }}
        .chart-image {{
            width: 100%;
            max-width: 100%;
            border: 1px solid #e5e7eb;
            border-radius: 8px;
        }}
        .footer {{
            margin-top: 3rem;
            padding-top: 1rem;
            border-top: 1px solid #e5e7eb;
            font-size: 8pt;
            color: #9ca3af;
            text-align: center;
        }}
        .table-section {{
            margin-bottom: 2rem;
            page-break-inside: avoid;
        }}
        .report-table {{
            width: 100%;
            border-collapse: collapse;
            font-size: 9pt;
        }}
        .report-table th {{
            background: #f3f4f6;
            border: 1px solid #e5e7eb;
            text-align: left;
            padding: 8px;
        }}
        .report-table td {{
            border: 1px solid #e5e7eb;
            padding: 8px;
        }}
    </style>
</head>
<body>
    <div class="header">
        <h1>{page_title_safe}</h1>
        <div class="subtitle">Relatório gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}</div>
    </div>
    
    <div class="filters">
        <strong>Filtros aplicados:</strong> {escape(filters_display)}
    </div>
"""

        # Add KPIs
        if kpis:
            html_content += '<div class="kpi-grid">'
            for kpi in kpis:
                html_content += f"""
    <div class="kpi-card">
        <div class="kpi-label">{escape(str(kpi.get('label', '')))}</div>
        <div class="kpi-value">{escape(str(kpi.get('value', '—')))}</div>
    </div>
"""
            html_content += '</div>'

        # Add charts
        for chart in charts:
            chart_title = escape(str(chart.get('title', 'Gráfico')))
            chart_image = str(chart.get('image', ''))
            html_content += f"""
    <div class="chart-section">
        <div class="chart-title">{chart_title}</div>
        <img class="chart-image" src="{chart_image}" alt="{chart_title}">
    </div>
"""

        for table in tables:
            columns = table.get('columns', [])
            rows = table.get('rows', [])
            html_content += f"""
    <div class="table-section">
        <div class="chart-title">{escape(str(table.get('title', 'Tabela')))}</div>
        <table class="report-table">
            <thead><tr>
"""
            for col in columns:
                html_content += f"<th>{escape(str(col))}</th>"
            html_content += "</tr></thead><tbody>"

            for row in rows[:MAX_TABELA_ROWS]:
                html_content += "<tr>"
                for cell in row:
                    html_content += f"<td>{escape(str(cell))}</td>"
                html_content += "</tr>"

            html_content += "</tbody></table></div>"

        html_content += """
    <div class="footer">
        NIR Dashboard — Sistema de Gestão de Leitos Hospitalares
    </div>
</body>
</html>
"""

        print("HTML gerado com sucesso", flush=True)
        
        # Generate PDF
        try:
            print("Iniciando conversão para PDF...", flush=True)
            pdf_bytes = HTML(string=html_content).write_pdf()
            print(f"PDF gerado: {len(pdf_bytes)} bytes", flush=True)
        except Exception as pdf_error:
            print(f"Erro na conversão WeasyPrint: {pdf_error}", flush=True)
            raise
        
        # Create response
        response = make_response(pdf_bytes)
        response.headers['Content-Type'] = 'application/pdf'
        filename = f"relatorio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        response.headers['Content-Disposition'] = f'attachment; filename={filename}'
        
        print("PDF enviado com sucesso", flush=True)
        return response

    except Exception as e:
        print(f"ERRO ao gerar PDF: {type(e).__name__}: {e}", flush=True)
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=80, debug=True)